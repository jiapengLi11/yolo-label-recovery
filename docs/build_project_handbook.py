from __future__ import annotations

import argparse
from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "YOLO_Label_Recovery_项目开发学习与大厂面试手册.docx"
ASSETS = ROOT / "docs" / "_handbook_assets"

NAVY = "17324D"
BLUE = "1F6F8B"
TEAL = "159A9C"
GOLD = "D99B30"
INK = "1A2530"
MUTED = "5E6C78"
LIGHT = "EFF5F5"
PALE_BLUE = "EAF2F6"
PALE_GOLD = "FFF5DA"
PALE_RED = "FCEBEC"
WHITE = "FFFFFF"
CODE_BG = "10212B"
CODE_FG = "E8F1F2"
GRID = "C9D7DC"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name: str, size: float, color: str = INK, bold: bool = False, italic: bool = False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(element, fill: str):
    props = element.get_or_add_pPr() if element.tag.endswith("}p") else element.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def paragraph_border(paragraph, color: str = TEAL, size: str = "12", side: str = "left"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "8")
    edge.set(qn("w:color"), color)
    borders.append(edge)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_image_alt(inline_shape, alt: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    link.append(run)
    paragraph._p.append(link)


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22

    title = doc.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color, before, after in (
        ("Heading 1", 17, NAVY, 16, 8),
        ("Heading 2", 13.5, BLUE, 12, 6),
        ("Heading 3", 11.5, TEAL, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Inches(0.38)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Cascadia Mono"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Cascadia Mono")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Cascadia Mono")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(7.7)
    code.font.color.rgb = rgb(CODE_FG)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.10)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    caption = doc.styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)


def add_header_footer(doc: Document):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("YOLO LABEL RECOVERY  |  工程学习与面试手册")
        set_run_font(r, "Microsoft YaHei", 8, MUTED, bold=True)
        paragraph_border(p, GRID, size="4", side="bottom")

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("内部学习版  ·  第 ")
        set_run_font(r, "Microsoft YaHei", 8, MUTED)
        add_page_field(p)
        r = p.add_run(" 页")
        set_run_font(r, "Microsoft YaHei", 8, MUTED)


def add_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None, italic: bool = False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, "Microsoft YaHei", 10.5, INK, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, "Microsoft YaHei", 10.5, INK, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, "Microsoft YaHei", 10.5, INK, italic=italic)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, "Microsoft YaHei", 10.2, INK)


def add_numbers(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r, "Microsoft YaHei", 10.2, INK)


def add_callout(doc: Document, title: str, text: str, tone: str = "info"):
    fill, accent = {
        "info": (PALE_BLUE, BLUE),
        "tip": (LIGHT, TEAL),
        "warn": (PALE_GOLD, GOLD),
        "risk": (PALE_RED, "B33A3A"),
    }[tone]
    p = doc.add_paragraph()
    shade(p._p, fill)
    paragraph_border(p, accent, size="18")
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.10)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title + "\n")
    set_run_font(r, "Microsoft YaHei", 10.4, accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, "Microsoft YaHei", 9.7, INK)
    return p


def add_formula(doc: Document, title: str, formula: str, explanation: str):
    """Render a readable formula block followed by its engineering meaning."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_font(r, "Microsoft YaHei", 10.2, BLUE, bold=True)
    p = doc.add_paragraph(style="Code Block")
    shade(p._p, "F4F8FA")
    paragraph_border(p, TEAL, size="12")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(formula)
    set_run_font(r, "Cambria Math", 11.5, NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    r = p.add_run(explanation)
    set_run_font(r, "Microsoft YaHei", 9.7, INK)


def add_logic_bridge(doc: Document, previous: str, current: str, following: str):
    add_callout(
        doc,
        "逻辑承接",
        f"前一步：{previous}\n本章解决：{current}\n下一步：{following}",
        "info",
    )


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    set_row_cant_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell._tc, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, "Microsoft YaHei", 9, WHITE, bold=True)
    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 0:
                shade(cell._tc, "F7FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, "Microsoft YaHei", 8.8, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def source_lines(relative: str, start: int, end: int) -> str:
    lines = (ROOT / relative).read_text(encoding="utf-8").splitlines()
    chosen = lines[start - 1 : end]
    return "\n".join(f"{idx:>4}  {line}" for idx, line in enumerate(chosen, start=start))


def xml_safe(text: str) -> str:
    """Remove control characters that OOXML cannot represent."""
    return "".join(char for char in text if char in "\t\n\r" or ord(char) >= 32)


def add_code(doc: Document, relative: str, start: int, end: int, note: str):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"代码清单  |  {relative}:{start}-{end}")
    set_run_font(r, "Microsoft YaHei", 8.7, BLUE, bold=True)
    p = doc.add_paragraph(style="Code Block")
    shade(p._p, CODE_BG)
    paragraph_border(p, TEAL, size="10")
    r = p.add_run(xml_safe(source_lines(relative, start, end)))
    set_run_font(r, "Cascadia Mono", 7.7, CODE_FG)
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = False
    r = p.add_run(note)
    set_run_font(r, "Microsoft YaHei", 8.5, MUTED, italic=True)


def add_command(doc: Document, command: str, note: str | None = None):
    p = doc.add_paragraph(style="Code Block")
    shade(p._p, CODE_BG)
    paragraph_border(p, GOLD, size="10")
    r = p.add_run(xml_safe(dedent(command).strip()))
    set_run_font(r, "Cascadia Mono", 8, CODE_FG)
    if note:
        p = doc.add_paragraph(style="Caption")
        r = p.add_run(note)
        set_run_font(r, "Microsoft YaHei", 8.5, MUTED)


def chapter(doc: Document, number: str, title: str, summary: str, *, new_page: bool = True):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = new_page
    r = p.add_run(f"{number}  {title}")
    set_run_font(r, "Microsoft YaHei", 17, NAVY, bold=True)
    add_callout(doc, "本章目标", summary, "tip")


def add_picture(doc: Document, path: Path, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_image_alt(shape, alt)
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    set_run_font(r, "Microsoft YaHei", 8.5, MUTED)


def pil_font(size: int, bold: bool = False):
    path = Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc")
    return ImageFont.truetype(str(path), size)


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color=BLUE, width=5):
    draw.line([start, end], fill="#" + color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 14 * sign, y2 - 9), (x2 - 14 * sign, y2 + 9)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 9, y2 - 14 * sign), (x2 + 9, y2 - 14 * sign)]
    draw.polygon(points, fill="#" + color)


def box(draw, xy, title, subtitle, fill, outline=BLUE):
    draw.rounded_rectangle(xy, radius=18, fill="#" + fill, outline="#" + outline, width=3)
    x1, y1, x2, _ = xy
    tf = pil_font(27, True)
    sf = pil_font(18)
    tb = draw.textbbox((0, 0), title, font=tf)
    draw.text(((x1 + x2 - (tb[2] - tb[0])) / 2, y1 + 24), title, font=tf, fill="#" + NAVY)
    lines = subtitle.split("\n")
    y = y1 + 72
    for line in lines:
        sb = draw.textbbox((0, 0), line, font=sf)
        draw.text(((x1 + x2 - (sb[2] - sb[0])) / 2, y), line, font=sf, fill="#" + MUTED)
        y += 27


def create_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)

    img = Image.new("RGB", (1500, 820), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "从模型证据到多人审核的数据闭环", font=pil_font(38, True), fill="#" + NAVY)
    nodes = [
        ((60, 170, 330, 330), "单类别 Teachers", "6 个模型串行加载\nFP16 + 流式推理", PALE_BLUE),
        ((420, 170, 690, 330), "审核证据", "GT/AUTO 四状态\nIoU + IoS + 中心距离", LIGHT),
        ((780, 170, 1050, 330), "CSV 导入桥", "流式读取\n每批不超过 500", PALE_GOLD),
        ((1140, 170, 1410, 330), "协作 API", "JWT + 项目隔离\n租约 + 版本号", PALE_BLUE),
        ((420, 500, 690, 660), "只读原始标签", "不在原目录写入\n保留可回滚性", "F5F5F5"),
        ((780, 500, 1050, 660), "人工决策", "新增 / 替换 / 拒绝\n全程审计", LIGHT),
        ((1140, 500, 1410, 660), "派生训练集", "决策完成后写回\n重新训练与评测", PALE_GOLD),
    ]
    for item in nodes:
        box(d, *item)
    for a, b in [((330, 250), (420, 250)), ((690, 250), (780, 250)), ((1050, 250), (1140, 250)), ((1275, 330), (1275, 500)), ((1140, 580), (1050, 580)), ((780, 580), (690, 580))]:
        draw_arrow(d, a, b)
    img.save(ASSETS / "architecture.png", quality=95)

    img = Image.new("RGB", (1450, 760), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "ReviewTask 状态机与保护机制", font=pil_font(38, True), fill="#" + NAVY)
    box(d, (70, 250, 340, 430), "PENDING", "公共待领队列\n按 ID 稳定排序", PALE_BLUE)
    box(d, (570, 250, 840, 430), "CLAIMED", "claimed_by + lease_until\n同一时刻仅一人拥有", LIGHT)
    box(d, (1070, 150, 1360, 330), "COMPLETED", "合法决策写入\n任务不可再次领取", PALE_GOLD)
    box(d, (1070, 450, 1360, 630), "ESCALATED", "UNCERTAIN 决策\n进入高级复核", PALE_RED)
    draw_arrow(d, (340, 330), (570, 330))
    draw_arrow(d, (840, 300), (1070, 240))
    draw_arrow(d, (840, 380), (1070, 530))
    draw_arrow(d, (570, 410), (340, 410), GOLD)
    d.text((360, 355), "释放 / 租约过期", font=pil_font(19, True), fill="#" + GOLD)
    d.text((390, 285), "SELECT ... FOR UPDATE", font=pil_font(18), fill="#" + MUTED)
    d.text((900, 205), "版本匹配", font=pil_font(18), fill="#" + MUTED)
    d.text((900, 500), "人工不确定", font=pil_font(18), fill="#" + MUTED)
    img.save(ASSETS / "task-state.png", quality=95)

    img = Image.new("RGB", (1500, 900), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((65, 40), "两名审核员同时领取任务时发生什么", font=pil_font(38, True), fill="#" + NAVY)
    cols = [(190, "Reviewer A"), (540, "Spring API"), (880, "MySQL"), (1230, "Reviewer B")]
    for x, name in cols:
        d.text((x - 70, 120), name, font=pil_font(23, True), fill="#" + BLUE)
        d.line((x, 165, x, 820), fill="#" + GRID, width=3)
    y = 230
    draw_arrow(d, (190, y), (540, y))
    d.text((250, y - 34), "claim-next(projectId)", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (540, y), (880, y))
    d.text((610, y - 34), "事务 + 行锁查询", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (880, y), (540, y))
    d.text((625, y - 34), "返回候选 R001，并锁定", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (1230, y), (540, y))
    d.text((940, y - 34), "几乎同时 claim-next", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (540, y), (880, y))
    d.text((620, y - 34), "R001 已 CLAIMED，选择 R002", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (540, y), (190, y))
    d.text((235, y - 34), "R001 + version=1 + lease", font=pil_font(18), fill="#" + INK)
    y += 90
    draw_arrow(d, (540, y), (1230, y))
    d.text((820, y - 34), "R002 + version=1 + lease", font=pil_font(18), fill="#" + INK)
    img.save(ASSETS / "claim-sequence.png", quality=95)

    img = Image.new("RGB", (1500, 780), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((65, 45), "推荐的亲手重建顺序", font=pil_font(38, True), fill="#" + NAVY)
    stages = [
        ("01", "问题定义", "漏标、联合场景\n不可变源数据"),
        ("02", "离线流水线", "Teacher 扫描\n阈值与审核包"),
        ("03", "数据库", "用户、项目\n任务、决策、审计"),
        ("04", "安全", "JWT、RBAC\n项目成员隔离"),
        ("05", "并发", "行锁、租约\n乐观版本"),
        ("06", "前端", "真实图像\n心跳与动作约束"),
        ("07", "交付", "Docker、测试\nCI 与面试材料"),
    ]
    for idx, (num, title, detail) in enumerate(stages):
        x = 55 + idx * 205
        d.ellipse((x, 215, x + 75, 290), fill="#" + TEAL)
        d.text((x + 18, 232), num, font=pil_font(22, True), fill="white")
        if idx < len(stages) - 1:
            draw_arrow(d, (x + 85, 252), (x + 190, 252), GRID, 4)
        d.text((x - 5, 330), title, font=pil_font(24, True), fill="#" + NAVY)
        y = 380
        for line in detail.split("\n"):
            d.text((x - 5, y), line, font=pil_font(18), fill="#" + MUTED)
            y += 30
    d.rounded_rectangle((65, 590, 1435, 700), radius=18, fill="#" + PALE_GOLD, outline="#" + GOLD, width=3)
    d.text((105, 620), "原则：每完成一层就写测试和演示，不要等到最后才验证。面试官更看重你的约束意识和排错过程。", font=pil_font(22, True), fill="#" + INK)
    img.save(ASSETS / "development-roadmap.png", quality=95)

    img = Image.new("RGB", (1500, 840), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((65, 40), "可信局域网多人审核的真实部署拓扑", font=pil_font(38, True), fill="#" + NAVY)
    box(d, (55, 180, 320, 350), "Reviewer A / B", "独立账号登录\n浏览器只访问 8088", PALE_BLUE)
    box(d, (430, 180, 760, 350), "Windows 审核主机", "Spring Boot standalone JAR\n内置 Vue 静态资源", LIGHT)
    box(d, (870, 110, 1180, 280), "只读审核包", "真实 review visuals\n源文件不被平台修改", PALE_GOLD)
    box(d, (870, 390, 1180, 560), "Ubuntu VMware", "MySQL 持久化\n任务、决策、审计", PALE_BLUE)
    box(d, (55, 560, 320, 730), "历史桌面审核", "company_decisions.csv\n4,465 条决定", PALE_GOLD)
    box(d, (430, 560, 760, 730), "幂等迁移接口", "按 candidate_id 对齐\n每批 <= 500", LIGHT)
    for start, end in [
        ((320, 265), (430, 265)),
        ((760, 235), (870, 195)),
        ((760, 300), (870, 475)),
        ((320, 645), (430, 645)),
        ((760, 645), (870, 525)),
    ]:
        draw_arrow(d, start, end)
    d.text((430, 390), "校园网 / 可信 LAN", font=pil_font(21, True), fill="#" + TEAL)
    d.text((430, 425), "JWT + 项目成员隔离 + 租约心跳", font=pil_font(18), fill="#" + MUTED)
    d.text((430, 455), "PESSIMISTIC_WRITE + @Version + 唯一键", font=pil_font(18), fill="#" + MUTED)
    img.save(ASSETS / "campus-deployment.png", quality=95)

    img = Image.new("RGB", (1500, 1040), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((65, 35), "一张联合场景图片如何走完整个数据闭环", font=pil_font(38, True), fill="#" + NAVY)
    stages = [
        ((45, 150, 315, 300), "01 源事实", "person + helmet + smoking\n原标签只有 helmet", PALE_RED),
        ((410, 150, 680, 300), "02 Teacher 证据", "person / smoking 有预测\nhelmet 与 GT 已匹配", PALE_BLUE),
        ((775, 150, 1045, 300), "03 几何分类", "同类已标 / 新目标\n跨类合理嵌套", LIGHT),
        ((1140, 150, 1410, 300), "04 审核队列", "稳定 candidate_id\n建议动作 + 可视化", PALE_GOLD),
        ((1140, 480, 1410, 630), "05 原子领取", "行锁 + 租约\n一个任务一个所有者", PALE_BLUE),
        ((775, 480, 1045, 630), "06 人工决定", "接受 / 替换 / 拒绝\n不确定则升级", LIGHT),
        ((410, 480, 680, 630), "07 安全写回", "重复复检 + 源 GT 校验\n只写派生 labels", PALE_GOLD),
        ((45, 480, 315, 630), "08 训练评测", "新版六类模型\n固定 test 对比", PALE_BLUE),
    ]
    for item in stages:
        box(d, *item)
    for start, end in [
        ((315, 225), (410, 225)), ((680, 225), (775, 225)), ((1045, 225), (1140, 225)),
        ((1275, 300), (1275, 480)), ((1140, 555), (1045, 555)), ((775, 555), (680, 555)), ((410, 555), (315, 555)),
    ]:
        draw_arrow(d, start, end)
    d.rounded_rectangle((75, 760, 1425, 935), radius=20, fill="#" + PALE_GOLD, outline="#" + GOLD, width=3)
    d.text((115, 790), "贯穿不变量", font=pil_font(28, True), fill="#" + NAVY)
    invariants = [
        "源图片和源标签只读", "Teacher 只提供证据", "候选身份稳定", "每任务只有一个最终决定",
        "失败可重试且不重复", "人工结论可追溯", "写回前再次校验", "固定测试集证明是否提升",
    ]
    for idx, text_value in enumerate(invariants):
        x = 120 + (idx % 4) * 325
        y = 845 + (idx // 4) * 42
        d.text((x, y), "• " + text_value, font=pil_font(19, idx < 4), fill="#" + INK)
    img.save(ASSETS / "end-to-end-lifecycle.png", quality=95)

    img = Image.new("RGB", (1600, 1120), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "矿区智能安全监控系统：在线处置与离线模型闭环", font=pil_font(38, True), fill="#" + NAVY)
    layers = [
        ((70, 150, 1530, 290), "设备与边缘层", "20 路 RTSP 摄像头  ·  Jetson / 推理节点  ·  设备心跳与断流重连", PALE_BLUE),
        ((70, 335, 1530, 475), "视频与推理层", "独立拉流  ·  最新帧槽  ·  有界队列  ·  动态 Batch  ·  frame_id / timestamp 对齐", LIGHT),
        ((70, 520, 1530, 660), "事件与业务层", "空间关联  ·  滑动窗口  ·  告警去重  ·  风险分级  ·  工单状态机", PALE_GOLD),
        ((70, 705, 1530, 845), "知识与智能层", "混合检索 RAG  ·  证据约束  ·  Function Calling  ·  人工审批", PALE_BLUE),
        ((70, 890, 1530, 1030), "数据与模型闭环", "困难样本回流  ·  Multi-Teacher  ·  多人审核  ·  派生数据集  ·  固定测试集评估", LIGHT),
    ]
    for item in layers:
        box(d, *item)
    for y1, y2 in ((290, 335), (475, 520), (660, 705), (845, 890)):
        draw_arrow(d, (800, y1), (800, y2), TEAL, 5)
    img.save(ASSETS / "mining-system-architecture.png", quality=95)

    img = Image.new("RGB", (1600, 900), "#F7FAFB")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "20 路视频：从最新帧到可靠告警", font=pil_font(38, True), fill="#" + NAVY)
    nodes = [
        ((50, 180, 300, 340), "独立 RTSP 通道", "断流隔离\n指数退避重连", PALE_BLUE),
        ((370, 180, 620, 340), "最新帧槽", "旧帧可覆盖\n不积压历史画面", LIGHT),
        ((690, 180, 940, 340), "共享有界队列", "背压与公平调度\n高风险通道升频", PALE_GOLD),
        ((1010, 180, 1260, 340), "GPU 动态 Batch", "数量/等待双触发\n单模型实例", PALE_BLUE),
        ((1330, 180, 1570, 340), "检测结果", "device + frame_id\ncaptured_at", LIGHT),
        ((1010, 560, 1260, 720), "时序事件", "人员空间关联\n20 秒多次命中", PALE_GOLD),
        ((690, 560, 940, 720), "业务告警", "迟滞 + 冷却\n唯一键去重", PALE_BLUE),
        ((370, 560, 620, 720), "工单闭环", "分级派发\n处理证据回填", LIGHT),
        ((50, 560, 300, 720), "困难样本回流", "误报 / 漏报\n进入补标平台", PALE_GOLD),
    ]
    for item in nodes:
        box(d, *item)
    for start, end in [
        ((300, 260), (370, 260)), ((620, 260), (690, 260)), ((940, 260), (1010, 260)),
        ((1260, 260), (1330, 260)), ((1450, 340), (1135, 560)), ((1010, 640), (940, 640)),
        ((690, 640), (620, 640)), ((370, 640), (300, 640)),
    ]:
        draw_arrow(d, start, end)
    img.save(ASSETS / "rtsp-event-pipeline.png", quality=95)


def build_document(output_path: Path = OUT):
    create_diagrams()
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("ENGINEERING FIELD GUIDE")
    set_run_font(r, "Cascadia Mono", 10, TEAL, bold=True)
    p = doc.add_paragraph(style="Title")
    r = p.add_run("YOLO Label Recovery\n项目开发学习与大厂面试手册")
    set_run_font(r, "Microsoft YaHei", 29, NAVY, bold=True)
    p = doc.add_paragraph()
    r = p.add_run("从多 Teacher 漏标恢复，到 Spring Boot + Vue 多人审核协作平台")
    set_run_font(r, "Microsoft YaHei", 14, BLUE, bold=True)
    p.paragraph_format.space_after = Pt(24)
    paragraph_border(p, TEAL, size="16", side="bottom")
    add_callout(
        doc,
        "这不是功能说明书",
        "本手册按真实工程问题和开发顺序组织。目标是让你能够不看稿画出架构、解释关键代码、复现并发保护、独立排查故障，并诚实回答系统当前边界。",
        "info",
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["本地仓库", r"E:\project11\yolo-label-recovery"],
            ["开源仓库", "github.com/jiapengLi11/yolo-label-recovery"],
            ["核心技术", "Python / YOLO / Spring Boot / Spring Security / JPA / MySQL / Vue / TypeScript / Docker / Windows / VMware"],
            ["学习目标", "能从需求、数据、并发、安全、测试和交付六个维度讲清项目"],
            ["版本基线", "中文面试作品集与代码学习路线，彻底更新日期 2026-08-30"],
        ],
        [2100, 7260],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("建议先通读，再按第 20 章动手重建。只有真正改过、跑过、故障过，项目才会成为你的能力。")
    set_run_font(r, "Microsoft YaHei", 11, MUTED, italic=True)

    doc.add_page_break()
    doc.add_heading("阅读方式与能力地图", level=1)
    add_paragraph(doc, "这份手册采用四层学习法：第一层补齐必要前置知识，第二层理解问题与架构，第三层沿真实代码追踪状态变化，第四层通过实验和故障复盘建立肌肉记忆。不要背代码；先说清输入、状态、不变量和失败分支，再解释实现。")
    add_table(
        doc,
        ["阶段", "建议投入", "你要获得的能力", "验收方式"],
        [
            ["通读", "4-6 小时", "知道各模块为什么存在", "10 分钟画出全链路"],
            ["代码跟读", "2-3 天", "能从 API 追到事务、实体和 SQL", "独立解释一次领取和决策"],
            ["动手重建", "5-7 天", "能修改功能并写测试", "新增一个角色或状态并通过 CI"],
            ["面试演练", "2 天", "能讲取舍、事故和边界", "30 分钟项目深挖不卡壳"],
        ],
        [1300, 1400, 3600, 3060],
    )
    add_callout(doc, "面试原则", "只把你能解释到代码和测试的内容说成“我实现了”。规划中的能力用“下一步会……”表达。", "warn")
    add_picture(doc, ASSETS / "development-roadmap.png", 6.6, "图 1  推荐的亲手重建顺序", "项目开发七阶段路线图")

    doc.add_heading("目录", level=1)
    toc = [
        "第零部分：前置知识与统一心智模型（第 0 章）",
        "第一部分：问题、需求与架构（第 1-4 章）",
        "第二部分：Python 数据治理与 Multi-Teacher（第 5-9 章）",
        "第三部分：多人协作平台后端与前端（第 10-16 章）",
        "第四部分：部署、测试、重建与面试（第 17-23 章）",
        "第五部分：真实部署、迁移与双账号联调复盘（第 24 章）",
        "第六部分：端到端案例推演与知识闭环（第 25 章）",
        "第七部分：用户反馈驱动的审核生产力迭代（第 26 章）",
        "第八部分：矿区智能监控系统与大厂面试技术篇（第 27-35 章）",
        "第九部分：GitHub 作品集、代码深挖与可复现实验（第 36-41 章）",
        "附录：API、命令、术语、代码索引与现场速查",
    ]
    add_numbers(doc, toc)

    chapter(doc, "0", "前置知识：先建立统一心智模型", "先掌握后文反复使用的检测监督、数据对象、HTTP、事务、并发与认证概念。")
    add_logic_bridge(doc, "你已经知道项目用了 YOLO、Python、Java 和 Vue", "把技术名词还原成数据、状态和约束", "带着统一术语进入真实业务问题")
    doc.add_heading("0.1 目标检测不是给图片分类", level=2)
    add_paragraph(doc, "图像分类回答‘整张图是什么’，目标检测回答‘图里有哪些目标、每个目标属于什么类、位于哪里’。YOLO 标签的一行通常由 class_id、中心点 x/y、宽 w、高 h 组成，坐标归一化到 0-1。训练时，模型不仅学习类别，还学习框位置；一个真实目标没有标签，并不等于它从损失中消失，它会落入背景监督并影响候选分配。")
    add_table(doc, ["概念", "输入/输出", "本项目为何关心"], [
        ["GT 标签", "人工确认的类别与框", "是训练监督和漏标判断基准，但并非天然完美"],
        ["Teacher 预测", "框、类别、置信度", "是发现疑似漏标的证据，不直接等于真值"],
        ["候选 Candidate", "预测 + 几何关系 + 建议动作", "是离线模型层与人工审核层之间的数据契约"],
        ["审核决策", "接受新增/替换/拒绝/不确定", "把模型证据转换成人工可追溯事实"],
        ["派生数据集", "源图像 + 原标签 + 已批准新增", "用于下一轮训练，且源标签始终可回滚"],
    ], [1800, 3300, 4260])

    doc.add_heading("0.2 漏标为什么会变成错误监督", level=2)
    add_paragraph(doc, "以‘戴安全帽并抽烟’为例：如果图片只标 helmet，没有标 smoking，模型在匹配正样本后，剩余候选位置会被当作背景或非 smoking。结果不是单纯‘少学一张’，而是同一视觉模式在别的图片中被要求输出 smoking、在这张图中又被压低。联合场景越多、漏标越系统，梯度冲突越明显。")
    add_formula(doc, "二分类交叉熵的直觉", "BCE(p, y) = -[ y log(p) + (1-y) log(1-p) ]", "真实 smoking 目标若漏标，相当于某些位置使用 y=0；当模型给出较高 p 时会受到惩罚。项目要修复的是监督语义，而不只是增加图片数量。")
    add_callout(doc, "重要边界", "检测器具体如何分配正负样本由版本和 TaskAlignedAssigner 等策略决定，但工程结论稳定：未标真实目标可能被当作负监督，完整一致的标签比盲目堆数据更重要。", "warn")

    doc.add_heading("0.3 Web 协作平台的最小知识", level=2)
    add_table(doc, ["术语", "一句话定义", "在本项目中的落点"], [
        ["HTTP 请求", "客户端携带方法、路径、头和正文访问服务", "Vue 通过 REST API 登录、领取、心跳和提交"],
        ["JWT", "服务端签名的身份与声明载体", "请求头 Bearer Token 表示当前账号和角色"],
        ["事务", "一组数据库动作要么全部成功，要么全部回滚", "领取、决策、审计必须保持一致"],
        ["悲观锁", "读数据时先阻止竞争者修改", "两个审核员同时领取时只能有一人拿到该行"],
        ["乐观版本", "提交时检查自己读取的版本是否仍最新", "拒绝旧页面覆盖别人已改变的任务"],
        ["租约", "有截止时间的临时所有权", "浏览器掉线后任务能够重新进入公共队列"],
        ["幂等", "重复请求不产生重复副作用", "任务和历史决策都可安全重跑导入"],
    ], [1700, 3400, 4260])

    doc.add_heading("0.4 全书只围绕六个问题", level=2)
    add_numbers(doc, [
        "输入是什么：源图片、原标签、Teacher 权重和策略配置从哪里来？",
        "状态是什么：候选、任务、租约、版本和决策分别处于什么阶段？",
        "不变量是什么：哪些事情无论成功失败都不能被破坏？",
        "正常路径是什么：数据怎样从离线扫描流向在线审核并回到训练集？",
        "失败路径是什么：OOM、重复请求、掉线、旧页面和非法路径怎样被阻断？",
        "证据是什么：用什么日志、测试、唯一键、统计和截图证明系统真的工作？",
    ])
    add_callout(doc, "阅读方法", "每读一段代码，都用这六个问题做注释。能回答它们，才算理解；只能复述类名和框架名，还不算掌握。", "tip")

    chapter(doc, "1", "业务问题：为什么不是再训练一次就结束", "把矿区联合场景漏标、错误监督和多人审核需求讲清楚。")
    add_logic_bridge(doc, "第 0 章建立了监督与状态的基本概念", "证明问题不是模型大小，而是训练事实不完整", "把风险翻译成可验证需求")
    add_paragraph(doc, "项目起点不是“做一个标注平台”，而是训练数据出现了系统性漏标：拖鞋图片只标拖鞋、吸烟图片只标吸烟，图中 person、helmet、vest 等有效目标可能被当作背景。多类别模型在这种数据上训练时，会收到互相矛盾的监督。")
    doc.add_heading("1.1 三个连续问题", level=2)
    add_numbers(doc, [
        "如何在 2 万级图片中找出可能漏标，而不是让人工从头翻一遍？",
        "如何区分“模型发现新目标”与“模型只是把同一个 GT 框画得更大或更小”？",
        "当 3-10 名审核员并行工作时，如何避免重复领取、旧页面覆盖和责任不清？",
    ])
    doc.add_heading("1.2 关键不变量", level=2)
    add_bullets(doc, [
        "源数据只读：任何自动化结果都写入派生目录，错误可以整批回滚。",
        "高置信度是证据，不是自动写标签的授权；联合场景和同目标歧义必须保留人工门控。",
        "每个任务最多形成一个最终决策；每个决策必须能追溯到审核人和时间。",
        "训练、审核、写回分阶段完成，不能把模型推理与人工决策混成一个不可审计脚本。",
    ])
    add_callout(doc, "你在面试中先讲什么", "先讲不完整标注如何形成错误监督，再讲系统如何保护数据和并发。不要一上来只报技术栈。", "tip")
    doc.add_heading("1.3 用一张联合场景图理解矛盾监督", level=2)
    add_table(doc, ["画面事实", "标签文件", "训练时可能学到什么", "正确修复"], [
        ["人 + 安全帽 + 吸烟", "只有 helmet", "helmet 是正样本；person、smoking 可能被当背景", "补 person、smoking，而不是删除 helmet"],
        ["人 + 拖鞋", "只有 slipper", "模型可能学会拖鞋附近的人体特征不是 person", "联合检查 person 与 slipper"],
        ["多人施工", "只标前景人员", "未标路人是否算 person 取决于标注规范", "先统一标注口径，再决定是否补"],
    ], [2200, 1900, 3200, 2060])
    add_paragraph(doc, "这里最容易犯的错误是把‘模型能检测到’等同于‘业务上应该标’。自动补标必须同时满足视觉证据和标注规范。比如远处模糊路人、海报人物、屏幕中的人是否属于 person，要由数据规范决定，不能只看置信度。")

    chapter(doc, "2", "需求工程：把想法变成可验证约束", "从用户故事、功能需求和非功能需求推导实现。")
    add_logic_bridge(doc, "上一章确认了错误监督与多人协作风险", "把模糊诉求变成状态、约束和验收条件", "用架构边界承载这些约束")
    doc.add_heading("2.1 角色与用户故事", level=2)
    add_table(doc, ["角色", "用户故事", "权限边界"], [
        ["ADMIN", "创建账号、项目、分配成员、导入任务", "可访问全部项目"],
        ["REVIEWER", "领取任务、查看真实图、提交或释放决策", "只能访问被分配项目"],
        ["AUDITOR", "查看进度与审计记录", "不能领取和提交决策"],
    ], [1500, 4700, 3160])
    doc.add_heading("2.2 非功能需求如何落到代码", level=2)
    add_table(doc, ["风险", "工程约束", "实现证据"], [
        ["两人拿到同一任务", "领取必须原子化", "PESSIMISTIC_WRITE + 事务"],
        ["浏览器停留后覆盖新结果", "提交必须携带读到的版本", "@Version + expectedVersion"],
        ["审核员掉线占坑", "所有权必须有期限", "lease_until + heartbeat"],
        ["重复导入 CSV", "导入操作可安全重试", "project_id + candidate_id 唯一键"],
        ["跨项目看图", "授权和路径都要校验", "membership + normalize/startsWith"],
        ["事故无法追责", "关键动作不可变记录", "audit_events"],
    ], [2000, 3000, 4360])
    add_callout(doc, "设计方法", "每个技术名词都应对应一个明确风险。反过来，如果你说不出风险，就不应该为了“显得高级”而引入技术。", "info")
    doc.add_heading("2.3 需求必须写成可失败的验收条件", level=2)
    add_table(doc, ["模糊说法", "可验证验收条件", "失败证据"], [
        ["支持多人审核", "两个账号并发领取时 candidate_id 不同", "拿到同一任务或出现重复决定"],
        ["不会丢进度", "进程中断后从最后已提交游标恢复", "重跑产生重复候选或跳过未处理图片"],
        ["数据安全", "源 labels 的全局哈希在前后保持一致", "源文件时间戳或哈希变化"],
        ["可追溯", "每个最终决定可查 reviewer、time、task 和 action", "只有最终 CSV，无法说明谁做的"],
    ], [1900, 4300, 3160])
    add_callout(doc, "工程化判断", "好的验收条件一定允许你回答‘怎样算失败’。只写‘性能好、稳定、可扩展’无法指导代码和测试。", "tip")

    chapter(doc, "3", "总体架构：离线 ML 与在线协作解耦", "理解模块边界、数据所有权和为什么没有把推理塞进 Java 服务。")
    add_logic_bridge(doc, "需求已经定义了源数据只读、可重放和多人原子领取", "为不同时间尺度和资源模型划分边界", "把边界落实到仓库和请求链")
    add_picture(doc, ASSETS / "architecture.png", 6.6, "图 2  从模型证据到多人审核的数据闭环", "YOLO 漏标恢复与多人审核总体架构")
    add_paragraph(doc, "Python 负责昂贵且可批处理的模型推理、几何判断、阈值分流和审核图生成；Web 平台负责账号、并发、权限、任务、决策与审计。两层通过 review_queue.csv 和受控导入 API 解耦。")
    doc.add_heading("3.1 为什么不让后端直接调用 GPU", level=2)
    add_bullets(doc, [
        "模型任务耗时长、资源波动大，与在线请求的低延迟和高可用目标冲突。",
        "离线流水线可以断点续跑、调整 batch、重新生成候选，而不会影响审核员登录。",
        "CSV 是可审计中间产物，便于交付公司、复核和重放；后续也可平滑替换成消息队列。",
    ])
    doc.add_heading("3.2 技术栈不是拼盘", level=2)
    add_table(doc, ["层", "技术", "选择理由"], [
        ["模型与数据", "Python / Ultralytics / Pillow", "生态成熟，适合 GPU 推理与图像处理"],
        ["在线 API", "Java 21 / Spring Boot", "事务、安全、JPA 和团队后端工程能力"],
        ["权限", "Spring Security / JWT", "无状态 API，角色声明进入授权链"],
        ["数据库", "MySQL / Flyway", "约束、事务、索引、可迁移 schema"],
        ["前端", "Vue 3 / TypeScript / Vite", "轻量、类型明确、快速构建审核工作台"],
        ["交付", "Docker Compose / Nginx", "统一启动、只暴露前端端口、同源代理"],
    ], [1300, 2800, 5260])
    doc.add_heading("3.3 数据所有权比语言边界更重要", level=2)
    add_table(doc, ["数据对象", "唯一写入者", "读取者", "为什么这样约束"], [
        ["源 images/labels", "原始数据治理流程", "Python 扫描与写回工具", "防止自动化悄悄污染基线"],
        ["候选证据 CSV", "Python Multi-Teacher", "桌面审核器、导入桥接", "保留可重放、可抽查的中间证据"],
        ["在线任务状态", "Spring Boot 事务服务", "Vue 与审计接口", "所有并发变化必须经过同一规则入口"],
        ["人工最终决策", "审核 API/桌面追加日志", "迁移、写回、报告", "人类结论与模型建议分离"],
        ["派生 labels", "review_apply.py", "下一轮训练与评测", "每轮都能回到源基线重新生成"],
    ], [1800, 2500, 2500, 2560])
    add_callout(doc, "边界判断", "Python/Java 只是实现语言；真正的边界是‘谁有权修改哪一类事实’。只要数据所有权清楚，未来可把 CSV 换成对象存储或消息队列而不破坏业务语义。", "info")

    chapter(doc, "4", "仓库结构与一次请求的完整路径", "能从目录定位职责，并从前端点击追踪到数据库。")
    add_logic_bridge(doc, "上一章确定离线证据与在线状态各归谁管理", "把一次用户动作映射到具体文件、方法和数据库变化", "进入离线候选生成算法")
    add_command(doc, """
    yolo-label-recovery/
      yolo_label_recovery/       # 数据审计、几何决策、审核与写回
      autolabel_with_single_class_models.py
      platform/
        backend/                 # Spring Boot REST API
        frontend/                # Vue 3 审核工作台
        tools/import_review_queue.py
        docker-compose.yml
      tests/                     # Python 单测与冒烟测试
      docs/                      # 架构、案例和本手册
    """, "目录不是按语言随意堆放，而是按离线能力、在线能力、测试和文档划分。")
    doc.add_heading("4.1 “点击接受新增”经过哪些层", level=2)
    add_numbers(doc, [
        "App.vue 根据 recommendedAction 计算当前允许按钮。",
        "api.ts 发送 decision、expectedVersion 和 comment，并自动携带 Bearer JWT。",
        "TaskController 取得当前用户，调用 TaskService.decide。",
        "TaskService 校验任务所有者、租约、版本和动作白名单。",
        "ReviewDecision 以 task_id 唯一约束落库，ReviewTask 进入 COMPLETED 或 ESCALATED。",
        "AuditService 追加 TASK_DECIDED 事件，前端刷新项目进度。",
    ])
    doc.add_heading("4.2 同一次点击的成功路径与失败路径", level=2)
    add_table(doc, ["检查点", "成功时", "失败时", "HTTP/数据结果"], [
        ["JWT 与项目成员", "识别当前用户且有项目权限", "令牌无效或未分配项目", "401/403，事务尚未开始业务修改"],
        ["任务所有权与租约", "任务归当前用户且租约有效", "任务属于别人或已过期", "409，过期任务可释放回队列"],
        ["expectedVersion", "页面版本等于数据库版本", "页面停留太久，版本已变化", "409，不覆盖新状态"],
        ["DecisionPolicy", "动作与 recommendation 相容", "把 replace 场景误提交为 add", "400，拒绝非法业务动作"],
        ["唯一决定与审计", "决定、任务状态、审计一起提交", "唯一键冲突或数据库异常", "整个事务回滚，不留下半条记录"],
    ], [1900, 2400, 2600, 2460])

    chapter(doc, "5", "Multi-Teacher：为什么每个类别一个模型", "理解 K×N 计算量、类别专长、单模型串行加载和数据流。", new_page=False)
    add_logic_bridge(doc, "在线链路已经能够安全消费候选", "解释候选怎样由六个单类模型生成", "用几何关系判断候选是否真是漏标")
    add_paragraph(doc, "多 Teacher 的核心不是模型数量，而是把“发现证据”和“最终真值”分开。每个单类别模型专注一种目标，扫描全量图片，输出未被同类 GT 覆盖的候选。对于 N 张图片和 K 个 Teacher，推理量约为 K×N，但显存中同时只保留一个模型。")
    add_code(doc, "autolabel_with_single_class_models.py", 754, 786, "外层按类别和数据划分迭代，内层按 current_batch 切片；model.predict 使用 stream=True 增量消费结果。")
    doc.add_heading("5.1 为什么要有负样本", level=2)
    add_paragraph(doc, "单类别模型如果只看正样本，会把其他类别或背景误认为目标。训练单类 Teacher 时保留适量其他场景作为空标签负样本，能够学习“不是该类”的边界；但负样本过多会稀释正样本，所以项目采用分类别策略，而不是统一比例。")
    add_callout(doc, "复杂度", "显存复杂度近似 O(1 个模型 + 1 个 batch)，CPU 内存不应随 K×N 的全部预测线性增长；磁盘承担流式 CSV 与派生标签。", "info")
    doc.add_heading("5.2 单类 Teacher 的训练集怎样构成", level=2)
    add_table(doc, ["样本类型", "标签处理", "作用", "过量风险"], [
        ["该类正样本", "仅保留目标类别并重映射为 class 0", "学习目标外观与框回归", "场景单一会过拟合"],
        ["其他类别场景", "空标签负样本", "学习相似背景和其他 PPE 不是目标", "太多会稀释正样本"],
        ["联合场景", "目标类必须完整标注", "学习 helmet+smoking、person+slipper 共现", "漏标会再次制造错误负监督"],
        ["困难负样本", "模型高分误检但人工确认不是目标", "压制最常见假阳性", "错误否定会伤害召回"],
    ], [1800, 2600, 2800, 2160])
    add_paragraph(doc, "Teacher 的职责是提高某一类别的候选发现能力，不是成为最终六类部署模型。单类训练后还要在独立 val/test 上检查 Precision、Recall、mAP 与典型联合场景，确认它适合作为‘证据生成器’。")
    doc.add_heading("5.3 为什么每个候选必须有稳定 candidate_id", level=2)
    add_paragraph(doc, "同一预测会流经 CSV、审核图、桌面决策、Web 任务和最终写回。如果只依赖文件行号，排序或断点续跑后身份就会变化。稳定 candidate_id 应由 split、图片相对路径、类别、框坐标和模型/策略版本等规范化字段生成，使重复导入、历史迁移和审计都指向同一候选。")

    chapter(doc, "6", "几何判断：IoU 不够时如何识别同一目标", "掌握 IoU、IoS、中心距离和面积比例的联合语义。")
    add_logic_bridge(doc, "Teacher 已经输出候选框", "判断候选是新目标、已标目标还是歧义框", "再结合置信度决定自动、复核或忽略")
    add_code(doc, "yolo_label_recovery/review_decision.py", 46, 72, "overlap_metrics 同时返回 IoU、IoS、候选覆盖率、GT 覆盖率、中心距离与面积比例；best_match 按多个指标选择最相关 GT。")
    doc.add_heading("6.1 四种 GT/AUTO 状态", level=2)
    add_table(doc, ["状态", "含义", "处理"], [
        ["GT0_AUTO0", "该图片/类别既没有 GT，也没有预测", "只统计，不生成候选"],
        ["GT1_AUTO0", "有 GT，Teacher 没预测到", "暴露模型漏检，不自动删除 GT"],
        ["GT0_AUTO1", "无 GT，Teacher 有预测", "疑似漏标，进入置信度和冲突判断"],
        ["GT1_AUTO1", "GT 与预测并存", "判断已标、同目标歧义、不同目标或跨类冲突"],
    ], [1500, 4200, 3660])
    doc.add_heading("6.2 为什么 IoS 有价值", level=2)
    add_paragraph(doc, "当 AUTO 框只覆盖 GT 的一小部分，或 GT 是人体大框而 AUTO 是局部框时，IoU 会因并集过大而偏低。IoS 用交集除以较小框面积，能识别“小框几乎被大框包含”。但它不能单独决定同一目标，还要结合中心距离和面积倍率。")
    add_callout(doc, "面试回答", "IoU 衡量总体重叠，IoS 识别包含关系，中心距离约束空间邻近，面积比例约束尺度差异。联合判断是为了降低“同一目标重复补框”和“合理跨类别嵌套被误杀”两类风险。", "tip")
    doc.add_heading("6.3 四个几何量怎样一起工作", level=2)
    add_formula(doc, "交并比 IoU", "IoU = intersection(A, B) / union(A, B)", "适合两个大小接近的框；越接近 1，整体位置和尺度越一致。")
    add_formula(doc, "小框覆盖率 IoS", "IoS = intersection(A, B) / min(area(A), area(B))", "适合包含关系；小框几乎完全落在大框内时，即使 IoU 较低，IoS 仍接近 1。")
    add_formula(doc, "归一化中心距离", "d_center = distance(centerA, centerB) / sqrt(area(smaller_box))", "用较小框尺度归一化后，不同分辨率和远近目标可比较；值越小，空间中心越接近。")
    add_formula(doc, "面积倍率", "r_area = max(areaA, areaB) / min(areaA, areaB)", "用于识别一个框是否比另一个大很多。倍率极大时，即使中心接近，也不能轻易视为同一标注框。")
    add_table(doc, ["示例", "IoU", "IoS", "中心距离", "解释"], [
        ["两个近似人框", "0.78", "0.91", "0.08", "大概率已标，避免重复新增"],
        ["人框内的安全帽框", "0.06", "0.96", "0.72", "跨类别合理包含，不应互相删除"],
        ["同一人的大框与半身框", "0.42", "0.88", "0.16", "同类包含式歧义，应替换或人工判断"],
        ["并排两个人", "0.02", "0.05", "1.45", "两个独立目标，可保留新候选"],
    ], [2300, 900, 900, 1300, 3960])
    add_callout(doc, "为什么不能只写一个总分", "几何量表达不同语义。把它们直接加权成单一分数虽然方便，却会掩盖‘高 IoS 但跨类别合理嵌套’等情况。本项目先分类关系，再给 recommendation，便于审核和解释。", "warn")

    chapter(doc, "7", "阈值策略：从拍脑袋到审核驱动校准", "理解分类别阈值、AUTO/REVIEW/IGNORE 三段式和 Wilson 下界。", new_page=False)
    add_logic_bridge(doc, "几何规则已经排除明显已标与冲突候选", "用置信度和审核数据控制自动化风险", "让流水线在有限显存和中断条件下稳定执行")
    add_code(doc, "yolo_label_recovery/review_policy.yaml", 1, 22, "当前六类别策略文件。不同类别的尺寸、数据量和视觉难度不同，因此阈值分开配置。")
    add_code(doc, "yolo_label_recovery/calibration.py", 110, 124, "Wilson 下界不是只看经验精度，而是在有限样本下给精度增加置信保守项。")
    add_paragraph(doc, "AUTO 阈值要求审核样本中精度的置信下界达标；REVIEW 阈值以保留正样本召回为目标。这样可把“自动补”的风险和“人工复核”的工作量分别控制。")
    add_callout(doc, "不能说错", "仓库的 Web 多人平台不会自行决定 AUTO 阈值；它消费 Python 已生成的候选和建议动作。阈值校准属于离线数据治理层。", "warn")
    doc.add_heading("7.1 Precision、Recall 与审核成本的关系", level=2)
    add_formula(doc, "精确率", "Precision = TP / (TP + FP)", "阈值越高通常 Precision 越高，自动补进去的错误更少；AUTO 区优先保护这个指标。")
    add_formula(doc, "召回率", "Recall = TP / (TP + FN)", "阈值过高会漏掉真实候选；REVIEW 区通过人工接住中等置信度候选，优先保护召回。")
    add_table(doc, ["分段", "目标", "典型动作", "风险控制"], [
        ["AUTO 高置信", "高 Precision", "可进入自动建议，但仍写派生目录", "用 Wilson 下界与抽样复核"],
        ["REVIEW 中置信", "保留 Recall", "全部进入人工审核", "按类别限制队列与优先级"],
        ["IGNORE 低置信", "控制噪声与成本", "保留统计，不生成正式补标", "抽样检查是否存在系统性漏检"],
    ], [1800, 2100, 2800, 2660])
    doc.add_heading("7.2 Wilson 下界为何比 9/10 更可信", level=2)
    add_paragraph(doc, "如果某阈值只审核 10 个样本，9 个正确得到 90%，但样本太少，真实精度不确定。Wilson 下界会同时考虑成功比例和样本量：同样 90%，90/100 的下界显著高于 9/10。只有下界达到业务要求，才允许把该阈值视为自动化候选。")
    add_formula(doc, "Wilson 置信下界（记忆语义即可）", "lower = adjusted_center - uncertainty_margin", "样本越少，uncertainty_margin 越大；样本越多且成功率稳定，下界才逐渐接近观察到的 Precision。")

    chapter(doc, "8", "内存、显存、OOM 与断点续跑", "能解释为什么图片不是一次全装内存，以及 OOM 后如何无重复重试。")
    add_logic_bridge(doc, "规则和阈值已经定义候选如何分类", "让六个 Teacher 扫全量时不因资源波动破坏结果", "把候选交给可解释的人机审核界面")
    add_code(doc, "autolabel_with_single_class_models.py", 768, 841, "一个 batch 成功推理后才形成候选并提交；发生 CUDA OOM 时丢弃尚未提交的批次，减半重试。")
    add_code(doc, "yolo_label_recovery/state.py", 14, 55, "运行签名防止错误参数续跑；状态采用临时文件后原子替换，避免中断留下半个 JSON。")
    add_table(doc, ["资源", "驻留内容", "控制手段"], [
        ["GPU 显存", "当前 Teacher、当前 batch 激活值和预测张量", "单模型串行、FP16、自适应 batch"],
        ["CPU 内存", "路径清单、当前批图片、同类 GT 缓存", "流式消费、不过度预取"],
        ["磁盘", "CSV、样本图、派生 labels、state/manifest", "增量写入、原子状态、幂等键"],
    ], [1800, 4100, 3460])
    add_callout(doc, "事故复盘模板", "先记录失败类别、split、batch、显存占用和已提交游标，再解释为什么重试不会重复写入。比只说“把 batch 调小”更能体现工程能力。", "tip")
    doc.add_heading("8.1 batch 变大时到底增加了什么", level=2)
    add_table(doc, ["内存来源", "随 batch 的变化", "释放时机", "常见误解"], [
        ["模型权重", "基本不随 batch 变化", "卸载当前 Teacher 后", "显存大部分不一定都是权重"],
        ["前向激活", "通常近似随 batch、分辨率增长", "该批推理完成后", "imgsz 从 640 到 832 增长不是线性的"],
        ["预测/NMS 张量", "随候选数量与 batch 增长", "结果消费并转 CPU 后", "小目标密集场景可能比普通图更占用"],
        ["CPU 解码图片", "随预取、workers、batch 增长", "生成器释放引用后", "图片路径列表不等于像素已全装内存"],
        ["CSV/可视化", "应流式落盘", "写入完成即可释放", "把所有候选留在 list 会导致内存线性增长"],
    ], [1900, 2100, 2100, 3260])
    add_formula(doc, "分辨率的粗略影响", "pixel_work proportional to batch x imgsz^2", "832 相比 640 的像素量约为 (832/640)^2 = 1.69 倍；因此显存充足也要以稳定吞吐和 OOM 恢复能力决定 batch。")
    doc.add_heading("8.2 为什么 OOM 重试像一个小事务", level=2)
    add_numbers(doc, [
        "读取当前游标，构造尚未提交的 batch。",
        "执行推理并在内存中形成临时候选，不立即推进状态。",
        "若成功，先写候选/标签，再原子更新 state 游标。",
        "若 CUDA OOM，丢弃临时候选、清缓存、batch 减半，从同一游标重试。",
        "恢复运行时校验 run signature，防止换了模型或阈值却续接旧状态。",
    ])
    add_callout(doc, "不变量", "任何游标只能在该批副作用完整持久化后前进。否则会出现‘状态说处理过，文件却没写完’的永久漏数据。", "risk")

    chapter(doc, "9", "桌面审核器：从候选级到按图片聚合", "理解为什么同一图片的多个候选要一起看，以及本地交付的价值。", new_page=False)
    add_logic_bridge(doc, "离线流水线已经生成带证据和建议动作的候选", "让人能在联合场景中安全改变候选状态", "把单人离线决策迁入有事务和权限的多人平台")
    add_picture(doc, ROOT / "docs/assets/grouped-review-preview.png", 6.45, "图 3  真实按图片聚合审核界面", "真实生产审核界面截图")
    add_paragraph(doc, "候选逐条看会丢失联合场景上下文。按图片聚合后，审核员可以同时看到 person、helmet、smoking 等框，识别跨类别合理嵌套、重复框和漏标。桌面 Tk 工具保留为离线、单人、拷贝即用方案；多人平台不是替代，而是解决协作新增问题。")
    add_bullets(doc, [
        "JSONL 追加日志先记录每次点击，CSV 快照使用原子替换。",
        "按 recommendation 禁用不合法操作，防止把 replace 场景误当 add。",
        "原图只加载一次，在图片内切换多个候选，提高审核速度。",
    ])
    doc.add_heading("9.1 四类人工动作不是四个随意按钮", level=2)
    add_table(doc, ["动作", "适用语义", "写回行为", "禁止场景"], [
        ["ACCEPT_ADD", "确认是原标签未覆盖的新目标", "追加候选框", "同目标歧义或已有同类 GT"],
        ["ACCEPT_REPLACE_GT", "原 GT 与候选指向同目标，但候选框更合理", "按引用替换指定 GT", "找不到原 GT 或源标签已变化"],
        ["ACCEPT_EVAL_LABEL", "人工确认评估集也应补标", "仅显式允许时写入 val/test", "默认不应污染固定评估基准"],
        ["REJECT", "误检、重复或不符合标注规范", "不改标签，保留拒绝证据", "不能用来悄悄删除原 GT"],
        ["UNCERTAIN", "证据不足或需高级复核", "保持未解决/转 ESCALATED", "不能直接进入训练集"],
    ], [2200, 2600, 2300, 2260])
    add_code(doc, "yolo_label_recovery/review_apply.py", 112, 195, "写回前再次拒绝空白/不确定决定、校验替换引用、检查源 GT 是否变化，并对新增框做重复复检。人工审核不是绕过安全规则的后门。")

    chapter(doc, "10", "数据库与领域模型：先让约束进入 Schema", "掌握六张核心表、唯一键、索引和为什么不能只靠 Java if。")
    add_logic_bridge(doc, "桌面端已经证明人工决策语义可用", "把任务、所有权、决定和审计变成可并发维护的数据模型", "用 Spring 分层和事务实现状态变化")
    add_code(doc, "platform/backend/src/main/resources/db/migration/V1__review_collaboration_schema.sql", 39, 75, "review_tasks 记录所有权、租约和版本；review_decisions 用 task_id 唯一键保证一任务一决策。")
    add_table(doc, ["表", "核心字段", "关键约束"], [
        ["app_users", "username/password_hash/role/enabled", "username 唯一"],
        ["review_projects", "review_root/status/created_by", "创建人外键"],
        ["project_members", "project_id/user_id/assigned_by", "项目+成员唯一"],
        ["review_tasks", "candidate/state/claimed_by/lease/version", "项目+候选唯一，领取索引"],
        ["review_decisions", "task/reviewer/decision/comment", "task_id 唯一"],
        ["audit_events", "actor/action/resource/details/time", "资源+时间索引"],
    ], [1900, 4200, 3260])
    add_callout(doc, "数据库约束的价值", "应用层可能有 bug，也可能有两个实例同时执行。唯一键是最后一道一致性防线；捕获冲突后，API 将重复记录计为 skipped，而不是生成脏数据。", "info")
    doc.add_heading("10.1 实体、约束与索引分别解决什么", level=2)
    add_table(doc, ["机制", "解决的问题", "本项目例子", "不能替代什么"], [
        ["实体关系", "表达用户、项目、任务、决策之间的语义", "ReviewTask belongs to ReviewProject", "不能阻止并发重复写"],
        ["NOT NULL/外键", "阻止残缺记录和悬空引用", "decision 必须指向真实 task", "不能保证业务动作合法"],
        ["唯一键", "保证全局只有一份最终事实", "project+candidate、decision.task_id", "不能告诉用户为何冲突"],
        ["普通/复合索引", "让筛选和排序避免全表扫描", "project+state+lease_until", "不能自动保证事务正确"],
        ["事务", "让多个写动作一起提交或回滚", "任务完成+决定+审计", "不能替代授权和参数校验"],
    ], [1500, 2800, 2700, 2360])
    add_paragraph(doc, "数据库设计顺序应从不变量开始，而不是先画表：一任务最多一个决定，所以 review_decisions.task_id 唯一；一个候选在同一项目只导入一次，所以 review_tasks(project_id, candidate_id) 唯一；审核员只能看被分配项目，所以 project_members 需要项目+用户唯一并被 Service 查询。")
    doc.add_heading("10.2 为什么领取查询需要复合索引", level=2)
    add_paragraph(doc, "claim-next 的过滤条件包含 project_id、state，以及 CLAIMED 状态下 lease_until 是否过期，并按 id 取最早一条。如果没有匹配访问模式的索引，数据库会扫描大量行后再锁定，竞争时间变长。索引不是越多越好；每个索引都会增加导入和状态更新成本，应通过 EXPLAIN 与慢查询验证。")

    chapter(doc, "11", "Spring Boot 分层：Controller、Service、Repository、Entity", "能够从请求入口追到事务边界，不把所有逻辑堆在 Controller。")
    add_logic_bridge(doc, "Schema 已经提供最后一道一致性防线", "让业务规则在可测试的事务服务中执行", "在服务入口建立身份和资源授权")
    add_table(doc, ["层", "在本项目中的职责", "禁止做什么"], [
        ["Controller", "参数接收、当前用户、HTTP 状态与 DTO", "不直接写业务并发逻辑"],
        ["Service", "授权、事务、状态迁移、审计", "不依赖前端是否正确"],
        ["Repository", "查询、锁语义、聚合统计", "不承载用户故事"],
        ["Entity", "字段、版本和合法状态变化", "不暴露任意 setter"],
    ], [1600, 4400, 3360])
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/api/TaskController.java", 28, 67, "Controller 很薄：解析请求并把当前用户传给 TaskService。")
    add_callout(doc, "你应该会画", "HTTP → SecurityFilterChain → Controller → Service(@Transactional) → Repository/JPA → MySQL；异常由 ResponseStatusException 转换为 HTTP 状态。", "tip")
    doc.add_heading("11.1 @Transactional 的真实含义", level=2)
    add_numbers(doc, [
        "Controller 调用 Spring 代理对象，而不是直接获得一个裸 Service。",
        "代理在方法前开启或加入数据库事务，并绑定当前线程的 EntityManager。",
        "Repository 查询得到受管理实体；修改字段后，JPA 可在 flush/commit 时生成 UPDATE。",
        "方法正常结束时提交；运行时异常传播时回滚决定、任务状态和审计写入。",
        "事务结束后实体不再受当前持久化上下文管理，不能依赖延迟加载给前端拼 DTO。",
    ])
    add_callout(doc, "常见陷阱", "同一个类内部 self 调用另一个 @Transactional 方法可能绕过代理；长时间在事务中做文件 I/O 或网络请求会延长锁持有。事务只包数据库一致性所需的最小范围。", "warn")
    doc.add_heading("11.2 为什么 Controller 薄、Service 厚、Entity 不贫血也不万能", level=2)
    add_paragraph(doc, "Controller 负责把 HTTP 世界转换为业务调用；Service 负责跨实体、授权、事务和审计；Entity 维护自身合法状态变化，例如 claim、release、extendLease。把所有逻辑塞 Entity 会让它依赖仓库和当前用户，把所有逻辑塞 Service 又会出现任意 setter。项目采用‘实体守局部状态，Service 守跨对象用例’。")

    chapter(doc, "12", "JWT、密码与 RBAC", "解释登录签发、请求验签、角色映射和项目级授权是两层不同问题。")
    add_logic_bridge(doc, "事务层能够安全改变状态", "确认是谁在调用、可否访问这个项目", "让批量导入和历史迁移具备安全边界")
    add_picture(doc, ROOT / "docs/assets/platform-login.png", 6.45, "图 6  真实角色登录界面", "真实平台登录页，展示角色边界与协作入口")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/config/SecurityConfig.java", 39, 58, "SecurityFilterChain 将登录与健康检查放行，其余 API 需要 JWT；管理员接口和任务写接口按角色限制。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TokenService.java", 20, 39, "TokenService 在 JWT 中写入 subject、role、签发和过期时间，使用 HS256 密钥签名。")
    add_paragraph(doc, "RBAC 只回答“这个角色是否能调用某类接口”；项目成员校验回答“这个用户是否能访问这个具体项目”。两者缺一不可。ADMIN 在业务服务中作为全局旁路，REVIEWER/AUDITOR 必须存在 project_members 记录。")
    add_callout(doc, "当前边界", "当前实现是短期访问令牌，没有 refresh token、撤销列表、SSO 或 MFA。生产化可接入企业 OIDC/Keycloak，并把密钥放入 Secret Manager。", "risk")
    doc.add_heading("12.1 Authentication 与 Authorization 不同", level=2)
    add_table(doc, ["阶段", "问题", "本项目机制", "失败结果"], [
        ["认证 Authentication", "你是谁？", "账号密码 + BCrypt，签发 JWT", "401"],
        ["角色授权 RBAC", "这个角色能调用哪类接口？", "ADMIN/REVIEWER/AUDITOR", "403"],
        ["资源授权", "你能访问哪一个项目/任务？", "project_members + task owner", "403/404"],
        ["业务约束", "当前状态允许这个动作吗？", "lease/version/DecisionPolicy", "400/409"],
    ], [1800, 2600, 3000, 1960])
    add_paragraph(doc, "只在前端隐藏按钮不算授权，因为攻击者可以直接构造 HTTP 请求。SecurityConfig 做接口级粗粒度限制，Service 的 requireAccess/requireOwnedTask 做资源级和状态级细粒度判断；两层缺一不可。")
    doc.add_heading("12.2 一次 JWT 请求的完整时序", level=2)
    add_numbers(doc, [
        "用户提交 username/password；服务端用 PasswordEncoder 比对哈希，不保存或返回明文。",
        "TokenService 使用服务端密钥签名 JWT，写入 subject、userId、role、issuedAt 和 expiresAt。",
        "Vue 把令牌放入 Authorization: Bearer <token>，api.ts 为后续请求统一添加。",
        "Resource Server 验证签名和过期时间，把 role 映射为 Spring Authority。",
        "Controller 从 Authentication 取得当前用户；Service 再查询账号是否 enabled 以及项目成员关系。",
    ])
    add_callout(doc, "安全理解", "JWT 只是防篡改的身份声明，不是加密保险箱。不要放密码和敏感业务数据；一旦泄露，在过期前可能被重放，所以必须使用 HTTPS、短有效期和安全存储。", "risk")

    chapter(doc, "13", "项目隔离与幂等批量导入", "掌握 bounded batch、streaming、唯一键和安全重试。")
    add_logic_bridge(doc, "账号和资源访问边界已经建立", "安全地把离线候选与历史决定灌入在线状态机", "处理多人同时领取和长时间审核")
    add_code(doc, "platform/tools/import_review_queue.py", 43, 80, "Python 生成器逐行读取 CSV，校验必需列，并把字段转换成 API DTO；不会把 3 万行全部构造成一个请求。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/ProjectService.java", 102, 139, "后端限制每批 1-500，阻止批内 candidateId 重复，并跳过数据库中已存在的候选。")
    add_paragraph(doc, "幂等性不是“同一请求永远返回一模一样”，而是重复执行不会重复产生副作用。导入中断后可以从头重放 CSV：已有 candidateId 被唯一约束识别，结果计入 skippedExisting。")
    doc.add_heading("13.1 历史桌面决定如何迁入", level=2)
    add_paragraph(doc, "多人 Web 平台上线时，桌面端已经完成了部分审核。正确做法不是重新审核，也不是直接改数据库，而是先导入完整任务，再按 candidate_id 对齐历史决定。服务端复用同一 DecisionPolicy 校验动作是否合法，并依赖 review_decisions.task_id 唯一键保证重复迁移不产生第二条决定。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/DecisionPolicy.java", 8, 25, "在线提交和历史迁移共享同一动作白名单，避免两条路径产生不同业务规则。")
    add_code(doc, "platform/tools/import_review_queue.py", 93, 113, "decision_rows 逐行读取桌面决定，只迁移已完成项，并把中文工具的动作值映射为后端枚举。")
    add_command(doc, """
    python platform/tools/import_review_queue.py review_queue.csv ^
      --base-url http://localhost:8088 ^
      --username admin ^
      --password "<your-password>" ^
      --project-name "2026-08 六类联合场景审核" ^
      --review-root /review-data ^
      --decisions company_decisions.csv ^
      --batch-size 500
    """, "先使用 --dry-run 同时校验任务和历史决定，再执行真实导入；重复运行只增加 skipped，不会重复写决定。")
    doc.add_heading("13.2 幂等不是‘代码先查一下’", level=2)
    add_paragraph(doc, "典型错误是先 SELECT 是否存在，不存在再 INSERT。两个并发请求可能同时查到不存在，随后都插入。正确做法是多层组合：批内 Set 去重减少无效请求，数据库唯一键提供原子防线，Service 捕获或统计冲突，API 返回 imported/skipped/unknown/invalid 让调用者知道发生了什么。")
    add_table(doc, ["层次", "幂等键", "重复时行为", "证据"], [
        ["任务 CSV 批内", "candidate_id", "同批重复直接拒绝", "dry-run 错误行"],
        ["任务数据库", "project_id + candidate_id", "已有任务计 skipped", "唯一键与导入统计"],
        ["历史决定", "task_id", "已有最终决定计 skippedExisting", "review_decisions 唯一键"],
        ["断点恢复", "run signature + class/split cursor", "从已提交边界继续", "state.json 与日志"],
    ], [1800, 2700, 2800, 2060])
    add_callout(doc, "重试语义", "重试安全不等于每次返回完全相同字符串，而是最终数据库副作用相同且重复项可解释。", "tip")

    chapter(doc, "14", "并发核心：悲观锁、租约与乐观版本", "能完整解释三种机制分别防什么，为什么需要组合。", new_page=False)
    add_logic_bridge(doc, "任务和决定可以安全重复导入", "定义多人竞争、掉线和旧页面提交时的所有权语义", "确保任何决策都合法、安全且可审计")
    add_picture(doc, ASSETS / "task-state.png", 6.5, "图 4  ReviewTask 状态机", "任务状态机及保护机制")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/repository/TaskRepository.java", 19, 31, "findClaimable 使用 PESSIMISTIC_WRITE，数据库在事务内锁住候选行；另一个事务不能同时领走同一行。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TaskService.java", 57, 85, "claimNext 先复用审核员尚未过期的当前任务，再查询一条可领取任务，写入所有者和租约后 flush。")
    add_picture(doc, ASSETS / "claim-sequence.png", 6.6, "图 5  两名审核员并发领取时序", "两名审核员同时领取任务的数据库时序")
    doc.add_heading("14.1 三种机制不能互相替代", level=2)
    add_table(doc, ["机制", "保护时刻", "防止的问题"], [
        ["悲观锁", "领取瞬间", "两事务同时选中同一 PENDING 行"],
        ["租约", "领取后到提交前", "审核员离线后永久占用任务"],
        ["乐观版本", "提交瞬间", "旧页面使用过期状态覆盖新状态"],
        ["决策唯一键", "数据库落库", "重复请求形成两条最终决策"],
    ], [1700, 2600, 5060])
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/domain/ReviewTask.java", 63, 73, "leaseUntil 表示临时所有权截止时间；@Version 由 JPA 自动递增并用于并发更新检测。")
    add_callout(doc, "典型追问", "为什么不用 Redis 分布式锁？当前任务真值已经在 MySQL，行锁和事务更直接且能与状态更新原子提交。规模扩大、领取热点明显时再评估 SKIP LOCKED、分片或队列。", "tip")
    doc.add_heading("14.2 两个审核员同时领取时数据库发生什么", level=2)
    add_numbers(doc, [
        "A 和 B 几乎同时开启事务，按 project/state/lease 查询第一条可领取任务。",
        "A 先获得该行 PESSIMISTIC_WRITE 锁；B 在同一行等待，不能读到并修改同一旧状态。",
        "A 将任务改为 CLAIMED、写 claimed_by/lease_until 并提交，行锁释放。",
        "B 重新评估查询条件时，该任务已不再可领取，于是锁定下一条任务或返回空。",
        "两人的事务最终返回不同 candidate_id；审计记录分别属于真实账号。",
    ])
    add_table(doc, ["场景", "只用悲观锁", "再加租约", "再加版本号"], [
        ["同时领取", "可以防重复", "同样可以", "不是主要作用"],
        ["领取后浏览器崩溃", "任务永久占用", "到期后可回收", "不能主动回收"],
        ["页面停留 20 分钟后提交", "领取锁早已释放", "只判断所有权是否过期", "拒绝旧版本覆盖"],
        ["重复提交决定", "锁不一定覆盖重试", "租约无关", "版本 + decision 唯一键共同阻断"],
    ], [2200, 2200, 2300, 2660])
    add_callout(doc, "隔离级别直觉", "行锁只在事务持有期间有效；浏览器审核持续几分钟，不可能一直持有数据库事务，所以必须把长期所有权建模成 lease，把旧读保护建模成 version。", "info")

    chapter(doc, "15", "决策约束、文件安全与审计", "理解服务端不信任前端、路径穿越防护和不可变审计。", new_page=False)
    add_logic_bridge(doc, "任务所有权和并发版本已经可靠", "阻止非法动作、越界读图并留下责任证据", "让前端正确维护任务与资源生命周期")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TaskService.java", 32, 36, "服务端动作白名单把推荐动作映射到允许决策；即使用户伪造请求也不能执行不合法动作。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TaskService.java", 109, 151, "decide 校验版本与动作；visual 对路径 normalize 后确认仍以项目根目录开头，阻止 ../ 越界读取。")
    add_paragraph(doc, "前端按钮禁用属于用户体验，后端 ALLOWED 校验才是安全边界。真实图片接口同样需要 JWT 和项目访问权，且容器将审核包挂载为只读目录。")
    add_callout(doc, "纵深防御", "权限校验防越权，路径归一化防目录穿越，只读挂载防写文件，审计表保留行为证据。四层分别解决不同攻击面。", "info")
    doc.add_heading("15.1 路径穿越为什么 normalize 后还要 startsWith", level=2)
    add_command(doc, r"""
    review_root = E:\review-package
    visual_file = ..\..\Windows\System32\config\SAM

    root   = Path(review_root).toAbsolutePath().normalize()
    visual = root.resolve(visual_file).normalize()
    allow only when visual.startsWith(root)
    """, "resolve 只是拼接路径；normalize 会消解 ..；startsWith(root) 才能确认最终绝对路径仍在项目审核根目录内。")
    add_paragraph(doc, "此外还要检查 Files.isRegularFile，防止目录或不存在路径被当作图片返回。生产环境应让审核包以只读权限挂载，并设置响应 Content-Type、缓存和大小限制；若允许符号链接，还要额外校验真实路径。")
    doc.add_heading("15.2 审计日志与普通业务日志不同", level=2)
    add_table(doc, ["记录", "目的", "应包含", "不应包含"], [
        ["业务日志", "排错和运行观测", "requestId、错误、耗时", "明文密码、完整 JWT"],
        ["审计事件", "证明谁在何时改变了什么", "actor、action、resource、details、time", "可随意 UPDATE 的最终状态"],
        ["模型证据", "复现候选来源", "model、policy、conf、box、GT 关系", "只有审核后的结论"],
    ], [1700, 2600, 3300, 1760])

    chapter(doc, "16", "Vue 审核工作台：状态、心跳与对象 URL", "能解释前端如何保持任务租约、加载受保护图片并约束决策。")
    add_logic_bridge(doc, "后端已提供安全、事务化的任务 API", "把异步请求、租约和图片资源组合成不易出错的用户流程", "将前后端一起打包部署并验证网络路径")
    add_picture(doc, ROOT / "docs/assets/platform-dashboard.png", 6.45, "图 7  真实多人审核进度看板", "真实运行平台的候选总量、待审核、协作占用、疑难升级和完成率")
    add_picture(doc, ROOT / "docs/assets/platform-review-multibox.png", 6.45, "图 8  真实五候选联合场景 Web 审核工作台", "同一地下场景聚合展示三个人员、一个背心和一个拖拉机候选，并完整呈现租约、置信度、快捷键和受约束决策")
    add_picture(doc, ROOT / "docs/assets/platform-admin.png", 6.45, "图 9  管理员账号与项目成员分配", "管理员创建独立审核账号并把账号分配到指定项目")
    add_code(doc, "platform/frontend/src/api.ts", 108, 137, "request 包装器集中处理 localStorage JWT、Authorization 头、204 和错误响应，并提供最近审核等强类型接口。")
    add_code(doc, "platform/frontend/src/App.vue", 63, 75, "decisionOptions 根据后端推荐动作动态生成按钮，但它只是体验层，不能替代服务端校验。")
    add_code(doc, "platform/frontend/src/App.vue", 381, 429, "受保护图片以 Blob 拉取并创建 object URL；心跳每 30 秒续租；切换任务时撤销旧 URL，防止内存泄漏。")
    add_bullets(doc, [
        "登录后 bootstrap 同时获取当前用户与可访问项目。",
        "选择项目后刷新进度；领取成功后立即加载真实审核图并启动心跳。",
        "提交成功清理任务、停止心跳、撤销 Blob URL，再刷新进度。",
        "AUDITOR 在 UI 中不能领取，但真正限制仍由后端角色规则完成。",
    ])
    add_callout(doc, "当前前端边界", "项目和用户数量较大时需要分页、搜索、虚拟列表；访问令牌放 localStorage 也需要结合 CSP/XSS 防护评估，或改为安全 Cookie。", "risk")
    doc.add_heading("16.1 前端需要维护哪些状态", level=2)
    add_table(doc, ["状态", "来源", "变化时机", "清理要求"], [
        ["accessToken/currentUser", "登录响应/bootstrap", "登录、刷新、退出", "401 或退出时清除"],
        ["selectedProject/progress", "项目与进度 API", "选项目、提交后刷新", "切项目前先释放活动任务"],
        ["task", "claim-next", "领取、提交、释放、租约过期", "任何离开任务路径都要置空"],
        ["visualUrl", "受保护 Blob 响应", "任务切换后重新创建", "revokeObjectURL 防内存泄漏"],
        ["heartbeatTimer", "setInterval", "领取后启动", "提交、释放、退出、组件卸载都停止"],
        ["busy/error", "异步动作", "请求开始/结束", "finally 中恢复，防按钮永久禁用"],
    ], [1900, 2200, 2500, 2760])
    doc.add_heading("16.2 为什么不能直接把受保护图片 URL 放进 img", level=2)
    add_paragraph(doc, "普通 img 标签无法方便地附加 Authorization 头。前端先用 fetch 携带 JWT 获取 Blob，再通过 URL.createObjectURL 生成本地临时 URL 给 img 使用。切换任务时必须 URL.revokeObjectURL，否则长期审核会让浏览器持有大量 Blob 内存。")
    add_callout(doc, "异步竞态", "如果用户快速切项目，旧请求可能晚于新请求返回。更完整的实现应使用 AbortController 或请求序号，只允许最新请求更新当前状态。", "warn")

    chapter(doc, "17", "Docker Compose、Nginx 与配置管理", "理解三服务启动、健康检查、同源代理和只读数据卷。")
    add_logic_bridge(doc, "前端与后端单独运行已正确", "把构建产物、配置、数据库和网络组织成可复现交付", "用测试和 CI 证明交付没有破坏语义")
    add_code(doc, "platform/docker-compose.yml", 1, 52, "MySQL 先通过健康检查，后端再启动；前端等待后端健康。审核包以 /review-data:ro 挂载。")
    add_code(doc, "platform/frontend/nginx.conf", 1, 19, "Nginx 同时托管 Vue 静态文件并代理 /api 和 /actuator，浏览器只访问一个源。")
    add_command(doc, r"""
    cd E:\project11\yolo-label-recovery\platform
    copy .env.example .env
    # 修改 .env 中数据库密码、Base64 JWT 密钥、管理员密码和审核包目录
    docker compose up --build -d
    docker compose ps
    docker compose logs -f backend
    """, "服务默认通过 http://localhost:8088 访问。不要把真实 .env 提交到 Git。")
    add_callout(doc, "生产差距", "单机 Compose 适合演示和小团队。生产环境还应加入 HTTPS、密钥托管、备份恢复、日志采集、指标告警、滚动升级和容量压测。", "warn")
    doc.add_heading("17.1 没有 Docker 时的 standalone JAR", level=2)
    add_paragraph(doc, "本轮真实联调采用 Windows 主机运行 Spring Boot，前端通过 Maven standalone profile 打入 JAR；MySQL 放在 Ubuntu VMware 中。这样浏览器只访问一个 8088 端口，同时保留关系数据库事务语义。campus.local.env 保存数据库、JWT、Java 和可选 VMware 路径，并由 Git 忽略。")
    add_command(doc, r"""
    cd platform\frontend
    npm ci
    npm run build

    cd ..\backend
    mvn -Pstandalone clean package

    cd ..\campus-deploy
    .\start-platform.ps1
    .\status-platform.ps1
    """, "start-platform.ps1 会检查 MySQL；数据库不可达时可无界面启动 VMware，随后启动 JAR 并轮询 /actuator/health。")
    add_code(doc, "platform/campus-deploy/start-platform.ps1", 39, 71, "启动器把数据库可达性作为前置条件，并把虚拟机启动与应用启动串成可复用流程。")
    doc.add_heading("17.2 浏览器访问平台时经过哪些网络层", level=2)
    add_numbers(doc, [
        "浏览器访问校园网主机 IP:8088，TCP 首先经过客户端网络、交换网络和 Windows 防火墙。",
        "standalone 模式由 Spring Boot 同时返回静态前端和 /api；Compose 模式可由 Nginx 统一入口反向代理。",
        "Spring Security 校验 JWT，Controller/Service 处理请求，JPA 通过连接池访问 VMware 中的 MySQL。",
        "读取审核图时，后端校验项目权限和安全路径，再从 Windows 只读审核包流式返回。",
        "任何一层失败都应分别检查：DNS/IP、端口、防火墙、健康接口、应用日志、数据库 TCP 和 SQL。",
    ])
    add_table(doc, ["症状", "优先检查", "不要先做什么"], [
        ["页面打不开", "端口监听、防火墙、同网段互通", "不要先重装前端依赖"],
        ["页面开但登录失败", "/actuator/health、API 日志、账号状态", "不要归因于浏览器缓存"],
        ["登录后无项目", "project_members 与当前账号", "不要直接改前端隐藏逻辑"],
        ["图片 404/403", "review_root、visual_file、路径校验、文件权限", "不要关闭安全校验"],
        ["偶发数据库错误", "VM 状态、连接池、MySQL 日志与慢查询", "不要无限增加重试"],
    ], [2000, 4300, 3060])

    chapter(doc, "18", "测试与 CI：证明系统而不是口头保证", "理解单元、事务集成、HTTP 端到端和多技术栈 CI。")
    add_logic_bridge(doc, "部署链路已经能够启动", "用分层测试证明正常路径、失败路径和跨技术栈契约", "把实现顺序整理成可亲手重建的路线")
    add_code(doc, "platform/backend/src/test/java/com/jiapeng/labelreview/api/ApiWorkflowIntegrationTests.java", 45, 92, "真实随机端口测试贯穿管理员登录、创建审核员、创建项目、分配成员、导入、领取、决策和进度。")
    add_paragraph(doc, "TaskServiceIntegrationTests 进一步验证重复领取被阻止、非法动作被拒绝、旧版本被拒绝、项目外成员被拒绝。Python 侧有 42 个测试和多组无 GPU 演示；Vue 使用 vue-tsc + Vite 构建做类型与生产打包校验。")
    add_code(doc, ".github/workflows/ci.yml", 1, 26, "Python job 安装依赖、静态检查、冒烟测试和 pytest，后续还执行可复现演示与隐私扫描。")
    add_code(doc, ".github/workflows/ci.yml", 86, 117, "Java 21 Maven 测试和 Node 24 Vue 构建是独立 job，任一失败都会让 CI 失败。")
    add_callout(doc, "真实调试故事", "首次 v1.1.0 远端 CI 中 Java、Vue 和 Python 测试均成功，但 GT/AUTO 演示仍断言旧的 4 条候选；本地复现确认实际为 5 条后，只更新测试夹具断言并重新推送，最终 main CI 全绿。", "tip")
    doc.add_heading("18.1 本轮新增的验证证据", level=2)
    add_code(doc, "platform/backend/src/test/java/com/jiapeng/labelreview/api/ApiWorkflowIntegrationTests.java", 109, 153, "HTTP 集成测试验证历史决定首次导入、未知候选统计、重复导入跳过，以及 COMPLETED/ESCALATED 进度。")
    add_paragraph(doc, "更新后验证结果为：Python 42/42、Java 6/6、Vue 生产构建通过、standalone JAR 打包成功；平台截图脚本使用临时 Chrome profile，拍完审核页后主动释放任务，不干扰同事。")
    doc.add_heading("18.2 测试金字塔在本项目如何分工", level=2)
    add_table(doc, ["层次", "速度/依赖", "重点验证", "代表案例"], [
        ["纯函数单测", "最快，无数据库/GPU", "几何分类、阈值、策略合法性", "高 IoS 跨类嵌套不被当重复"],
        ["组件测试", "快，文件或内存对象", "CSV 解析、断点状态、写回安全", "不确定决策禁止生成派生集"],
        ["Repository/事务测试", "中等，真实数据库语义", "锁、唯一键、版本与回滚", "两事务不能领取同一任务"],
        ["HTTP 集成测试", "较慢，启动 Spring 上下文", "认证、授权、DTO、状态码和完整工作流", "导入→分配→领取→决定→进度"],
        ["浏览器/现场验收", "最慢，真实网络和文件", "构建产物、交互、截图、双账号", "校园网两台电脑同时审核"],
    ], [1900, 2200, 3100, 2160])
    add_callout(doc, "测试设计", "每个成功路径至少配一个相邻失败路径：能领取，也要测未授权；能提交，也要测过期租约、旧版本和非法动作；能迁移，也要测重复迁移和未知候选。", "tip")

    chapter(doc, "19", "从零开发顺序：你应该怎样亲手重建", "给出可以照着执行的工程顺序，每一步都有可验证产物。")
    add_picture(doc, ASSETS / "development-roadmap.png", 6.6, "图 10  从零重建七阶段", "从需求到交付的开发顺序")
    steps = [
        ("Step 1  写问题定义", "用一页文档固定漏标、联合场景、源数据只读和多人协作风险。产物：ADR 与验收表。"),
        ("Step 2  做无模型数据审计", "先检查标签格式、损坏图片、重复与数据泄漏。产物：dataset_audit.json/html。"),
        ("Step 3  实现单 Teacher dry-run", "只输出 candidates_all.csv，不写标签；加入同类 GT 匹配和阈值分流。"),
        ("Step 4  加入流式与恢复", "单模型串行、stream=True、批次提交、OOM 减半、state 原子写。"),
        ("Step 5  建人工审核包", "枚举 GT/AUTO 四状态，渲染同图多框，记录显式决策，不修改源标签。"),
        ("Step 6  设计数据库", "先写 Flyway V1 和唯一键，再写 JPA Entity；用 H2 测试迁移和约束。"),
        ("Step 7  建认证与项目隔离", "实现登录/JWT/RBAC，然后补 project_members 业务授权。"),
        ("Step 8  实现领取事务", "先写并发测试，再实现行锁、租约、心跳、释放和 @Version。"),
        ("Step 9  实现 Vue 工作台", "从 api.ts 类型与请求封装开始，再做登录、项目、图片、决策和管理抽屉。"),
        ("Step 10  打通 CSV 与历史决定迁移", "先 dry-run schema 校验，再实现有界任务批次、DecisionPolicy 复用和幂等重试。"),
        ("Step 11  完成两种交付", "先用 Compose 串起 MySQL/API/Nginx，再补 standalone JAR 与可信局域网启动脚本。"),
        ("Step 12  CI 与作品集", "每层加入自动测试、真实截图、架构说明、边界和故障复盘。"),
    ]
    for title, body in steps:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        set_run_font(r, "Microsoft YaHei", 11, BLUE, bold=True)
        add_paragraph(doc, body)

    chapter(doc, "20", "动手实验：把知识变成自己的", "提供从 30 分钟到 2 天的实验，要求你修改代码并留下测试证据。")
    labs = [
        ["Lab 1", "30 分钟", "画架构并口述一次决策链路", "不看文档能说出 6 层调用"],
        ["Lab 2", "1 小时", "把租约从 5 分钟改为 2 分钟并写过期测试", "过期任务可重新领取"],
        ["Lab 3", "2 小时", "新增 SUPERVISOR 角色，只能处理 ESCALATED", "权限测试和 UI 按钮通过"],
        ["Lab 4", "2 小时", "为 projects 增加分页和名称搜索", "大数据量不一次全返回"],
        ["Lab 5", "半天", "模拟两个线程同时 claim-next", "任务 ID 不重复"],
        ["Lab 6", "半天", "导入同一 CSV 两次", "第二次 created=0、skipped=N"],
        ["Lab 7", "1 天", "把 JWT 改成企业 OIDC 或 Keycloak", "移除自签密码登录依赖"],
        ["Lab 8", "1 天", "加入审核吞吐、租约过期、冲突次数指标", "Prometheus 可抓取"],
        ["Lab 9", "半天", "迁移同一份历史决定两次", "首次 imported=N，第二次 imported=0"],
        ["Lab 10", "半天", "两个独立账号同时领取并核对审计", "任务 ID 不同且 reviewer 可追溯"],
    ]
    add_table(doc, ["实验", "时间", "任务", "验收"], labs, [1100, 1200, 4200, 2860])
    add_callout(doc, "学习记录", "每个实验保留：需求、设计、关键 diff、失败日志、测试结果和反思。面试时最有价值的是这些真实证据。", "tip")
    doc.add_heading("20.1 第一次本地启动建议", level=2)
    add_command(doc, r"""
    cd E:\project11\yolo-label-recovery
    conda run -n yolo26 python -m pip install -e ".[dev]"
    conda run -n yolo26 pytest

    cd platform\backend
    mvn test

    cd ..\frontend
    npm ci
    npm run build
    """, "如果要启动完整多人平台，需要安装 Docker Desktop 并准备 .env；只学习代码可先跑三套测试。")

    chapter(doc, "21", "故障排查：从症状到证据", "建立分层排错方法，而不是靠重复重启。")
    add_table(doc, ["症状", "先看什么", "高概率原因", "验证动作"], [
        ["登录 401", "后端日志/JWT 时钟", "密钥不一致、令牌过期、用户禁用", "解码 claims，检查 env"],
        ["领取 409", "当前用户 active task", "跨项目仍有有效租约", "释放旧任务或等过期"],
        ["图片 404/403", "visual_file 与 review_root", "文件缺失或路径越界", "容器内 ls + normalize"],
        ["重复审核", "task_id/candidate_id/版本", "绕过领取、旧客户端、唯一键缺失", "并发测试与数据库约束"],
        ["导入内存高", "请求体与 Python 列表", "一次加载全 CSV", "生成器 + 500 条批次"],
        ["GPU OOM", "类别/split/batch", "激活值过大或模型未释放", "自适应减半并检查稳定 batch"],
        ["前端框延迟", "frame_id/时间戳", "检测坐标画在不同视频帧", "后端同帧画框或帧缓存对齐"],
    ], [1800, 2300, 3100, 2160])
    add_callout(doc, "排错顺序", "先缩小到层：浏览器、Nginx、API、安全、事务、数据库、文件系统或 GPU；再构造最小复现；最后改代码并补回归测试。", "info")

    chapter(doc, "22", "大厂面试讲法：从 90 秒到 30 分钟", "形成真实、可深挖、不过度承诺的项目叙事。")
    doc.add_heading("22.1 90 秒版本", level=2)
    add_callout(doc, "示范回答", "我负责的项目源于矿区六类检测数据的联合场景漏标。第一阶段我把每类单模型作为 Teacher，采用单模型串行、FP16 和流式批处理扫描 2 万级数据，用 IoU、IoS、中心距离和面积比区分已标与疑似漏标，并通过分类别阈值和人工门控生成派生标签。第二阶段为解决多人并行审核的重复领取和追溯问题，我实现了 Spring Boot + Vue 协作平台：JWT/RBAC 与项目成员隔离保证权限，MySQL 行锁保证原子领任务，租约处理掉线，JPA 版本号阻止旧页面覆盖，唯一键和审计表保证幂等与可追溯。平台实际导入 30,183 条候选，幂等迁移 4,465 条桌面决定，并用两个独立账号完成校园网并发审核验证。项目有 Python 42 项、Java 6 项测试、Vue 构建和 standalone JAR 证据。", "info")
    doc.add_heading("22.2 高频追问与回答骨架", level=2)
    qa = [
        ["为什么拆 Python 和 Java？", "推理是离线、高资源、可重放任务；协作是在线、事务和权限问题。CSV/API 解耦后可独立扩缩。"],
        ["行锁会不会慢？", "当前领取一次锁一行且有复合索引，规模适配小团队；热点增大可考虑 SKIP LOCKED、队列或分片。"],
        ["为什么既有悲观锁又有版本号？", "前者保护领取竞争，后者保护长时间审核后的旧提交，作用时刻不同。"],
        ["如何证明没重复补标签？", "同类 GT 几何匹配、候选去重、candidate_id 唯一、决策唯一、写回前再检查，分层防护。"],
        ["自动补标为什么不全自动？", "Teacher 预测不是事实，联合场景存在同目标尺度歧义和跨类嵌套；高风险类别必须保留人工证据。"],
        ["最大技术难点？", "不是写界面，而是定义任务所有权和一致性：领取、掉线、旧页面、重复请求和可追溯必须同时成立。"],
        ["如果上生产还缺什么？", "企业身份、HTTPS/密钥托管、可观测性、备份、分页、压测、灾备、内容安全与数据生命周期策略。"],
    ]
    add_table(doc, ["问题", "回答骨架"], qa, [2500, 6860])
    doc.add_heading("22.3 白板时必须画出的四样东西", level=2)
    add_bullets(doc, [
        "离线 Teacher → review_queue.csv → 在线 API → 人工决策 → 派生训练集。",
        "ReviewTask 的 PENDING / CLAIMED / COMPLETED / ESCALATED 状态机。",
        "两名审核员并发领取时的数据库行锁时序。",
        "JWT 角色授权与 project_members 资源授权的双层边界。",
    ])
    add_callout(doc, "不要这样说", "不要声称已完成 Kubernetes、Redis、TensorRT 20 路生产压测、SSO 或亿级任务调度，除非你补齐代码和证据。可以把它们作为演进方案。", "risk")

    chapter(doc, "23", "生产化路线图与技术债", "学会识别当前方案的适用范围，并给出演进而非推倒重来。")
    roadmap = [
        ["P0 安全", "HTTPS、Secret Manager、OIDC、登录限流、CSP", "上线前"],
        ["P0 数据", "MySQL 备份恢复、审核包校验和、数据保留策略", "上线前"],
        ["P1 可观测", "Micrometer/Prometheus、结构化日志、审计检索", "试运行"],
        ["P1 性能", "项目/用户分页、领取压测、连接池和索引分析", "任务量增长"],
        ["P2 协作", "高级复核队列、任务批次、评论和导出", "团队扩大"],
        ["P2 ML 闭环", "审核结果回流阈值校准、主动学习、漂移监控", "稳定运营"],
    ]
    add_table(doc, ["优先级", "改进", "触发时机"], roadmap, [1500, 5500, 2360])
    add_paragraph(doc, "架构演进应保留已验证的不变量：源数据不可变、决策可审计、任务只有一个最终结果、权限按项目隔离。替换基础设施时不要破坏这些语义。")

    chapter(doc, "24", "真实落地复盘：从桌面 CSV 到校园网多人审核", "按实际开发顺序复盘部署、迁移、双账号验证、故障修复与面试证据。")
    add_picture(doc, ASSETS / "campus-deployment.png", 6.6, "图 11  Windows + VMware + 校园网真实部署拓扑", "两名审核员通过可信局域网访问 Windows standalone JAR，MySQL 运行在 Ubuntu VMware，审核图片只读")
    add_callout(doc, "本轮可验证结果", "平台导入 30,183 条审核候选，迁移 4,465 条桌面历史决定；迁移时形成 4,464 条 COMPLETED 和 1 条 ESCALATED。两个独立账号已同时领取和审核，未观察到重复任务或覆盖问题。", "info")

    doc.add_heading("24.1 为什么选择这套部署，而不是立即上云", level=2)
    add_paragraph(doc, "参与审核的电脑处于同一可信校园网，任务规模属于小团队协作，真实审核图又占用较大磁盘。把应用放在已有 Windows 主机、数据库放在 Ubuntu VMware，可以复用本机数据和中间件，减少复制与云成本，同时保留 MySQL 的事务、唯一键和审计能力。")
    add_table(doc, ["组件", "本轮位置", "为什么这样放", "边界"], [
        ["Vue 前端", "Spring Boot JAR 内", "同源访问，不额外维护 Vite/Nginx", "仍需浏览器兼容与 CSP"],
        ["Spring Boot API", "Windows 主机:8088", "直接访问只读审核包，启动简单", "主机休眠后不可用"],
        ["MySQL", "Ubuntu VMware", "复用稳定数据库环境和事务能力", "需处理 VM 启动与备份"],
        ["审核图片", "Windows 本地目录", "避免复制数 GB 图片", "路径必须只读且防穿越"],
        ["审核客户端", "校园网浏览器", "无需安装开发环境", "跨网络需 VPN/TLS"],
    ], [1500, 1800, 3700, 2360])
    add_callout(doc, "架构判断", "这不是“Windows 比 Docker 更好”，而是根据网络、数据位置和团队规模选择交付形态。Docker Compose 仍是可移植基线，standalone JAR 是本轮现场约束下的最小可用方案。", "tip")

    doc.add_heading("24.2 从源码到可访问平台发生了什么", level=2)
    add_numbers(doc, [
        "Vue 执行 npm ci 与 npm run build，生成 frontend/dist。",
        "Maven 使用 standalone profile，把 dist 复制到 classpath/static，再打成可执行 JAR。",
        "campus.local.env 注入数据库、JWT、管理员、Java 与可选 VMware 路径；该文件不进 Git。",
        "start-platform.ps1 先测试 MySQL TCP 端口；不可达时用 vmrun 无界面启动 Ubuntu。",
        "脚本启动 java -jar，轮询 /actuator/health，成功后输出可配置的局域网地址。",
        "防火墙脚本只开放 TCP 8088，并允许显式传入校园网 CIDR，而不是写死个人 IP。",
    ])
    add_code(doc, "platform/backend/pom.xml", 86, 104, "standalone profile 将前端 dist 作为 Spring Boot 静态资源打包，实现一个进程交付。")
    add_code(doc, "platform/campus-deploy/start-platform.ps1", 73, 101, "启动器把日志、PID、Java 进程和健康检查统一管理，避免手工开多个终端。")
    add_command(doc, r"""
    cd E:\project11\yolo-label-recovery\platform\frontend
    npm ci
    npm run build

    cd ..\backend
    mvn -Pstandalone clean package

    cd ..\campus-deploy
    .\start-platform.ps1
    .\status-platform.ps1
    """, "真实密码、JWT、数据库地址和 VM 路径只写入 campus.local.env，不写入命令历史、截图或仓库。")

    doc.add_heading("24.3 历史审核迁移不是复制 CSV", level=2)
    add_paragraph(doc, "桌面 CSV 是离线事实记录，Web 数据库是在线任务状态。迁移必须先验证 candidate_id 与完整任务队列一一对应，再把动作映射到 DecisionType；未知候选、非法动作和已存在决定分别计数，不能悄悄丢弃或覆盖。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/ProjectService.java", 152, 215, "历史迁移在一个事务内完成候选查找、已有决定过滤、动作校验、任务状态迁移、决定保存和审计记录。")
    add_table(doc, ["输入情况", "服务端处理", "原因"], [
        ["candidate_id 不存在", "unknownCandidates +1，不写入", "防止决定挂错任务"],
        ["已有最终决定", "skippedExisting +1", "重复运行保持幂等"],
        ["动作不符合 recommendation", "invalidDecisions +1", "迁移不能绕过业务规则"],
        ["UNCERTAIN", "任务转 ESCALATED", "保留高级复核语义"],
        ["合法接受或拒绝", "任务转 COMPLETED", "恢复桌面审核进度"],
    ], [2300, 3200, 3860])
    add_callout(doc, "关键区别", "任务导入的幂等键是 project_id + candidate_id；决定导入的最终防线是 review_decisions.task_id 唯一键。二者解决不同层次的重复。", "tip")

    doc.add_heading("24.4 双账号并发验证怎么做才算有证据", level=2)
    add_numbers(doc, [
        "管理员创建两个独立 REVIEWER 账号，分别分配到同一项目，禁止共享管理员账号。",
        "两台浏览器或两台电脑同时登录，并在相近时间点击“领取下一条任务”。",
        "记录两边 candidate_id，确认不同；进度看板的“协作占用”随租约变化。",
        "一方提交决定，另一方继续心跳和提交，确认没有 409、重复决定或任务串号。",
        "关闭其中一个页面，等待租约到期或主动释放，确认任务重新回到公共队列。",
        "检查 audit_events 和 decision reviewer，确认操作归属真实账号。",
    ])
    add_table(doc, ["机制", "验证动作", "预期证据"], [
        ["悲观锁领取", "两个账号同时 claim-next", "返回不同 task/candidate"],
        ["租约", "页面持续打开与关闭", "heartbeat 延期；关闭后可回收"],
        ["乐观版本", "使用旧 expectedVersion 提交", "返回 409，不覆盖新状态"],
        ["唯一决定", "重复提交或迁移", "数据库拒绝第二条，接口计 skipped"],
        ["项目隔离", "未分配账号访问项目", "403 或不可见"],
    ], [1900, 3300, 4160])
    add_callout(doc, "诚实边界", "双账号实测证明基础分配和协作链路可用，不等同于 100 人并发压测。最大吞吐仍需 Gatling/JMeter、连接池监控和慢查询证据。", "warn")

    doc.add_heading("24.5 本轮真实故障与修复", level=2)
    add_table(doc, ["现象", "根因", "修复", "形成的工程经验"], [
        ["Spring Boot 启动时 Flyway 不完整", "Spring Boot 4 将集成拆为模块", "使用 starter-flyway + flyway-mysql", "升级依赖先查自动配置边界"],
        ["JAR 启动但没有前端", "默认 resources 不包含 dist", "增加 standalone Maven profile", "部署产物必须做端到端打开验证"],
        ["同事无法稳定访问", "局域网地址、防火墙和客户端隔离需区分", "监听 0.0.0.0，网段化防火墙，健康检查", "先分层判断网络还是应用"],
        ["历史审核散落在另一台电脑", "CSV 与数据库状态未对齐", "候选校验、归档、幂等迁移", "数据迁移必须有核验报告"],
        ["登录截图只有背景", "页面动画尚未完成", "等待 login-card opacity=1 再截图", "自动化要等待业务就绪而非只等 DOM"],
        ["默认 Python 无 pytest", "PATH 指向 base 环境", "使用 conda run -n yolo26", "环境证据要写进复现命令"],
        ["推送前远端已有新提交", "另一轮文档更新先落到 main", "fetch + rebase 后普通 push", "不强推覆盖协作者历史"],
    ], [1650, 2200, 2600, 2910])

    doc.add_heading("24.6 建议的代码阅读顺序", level=2)
    add_table(doc, ["顺序", "文件", "阅读问题"], [
        ["1", "import_review_queue.py", "CSV 如何流式校验、分批、重试和迁移决定？"],
        ["2", "ProjectController.java", "为什么历史迁移仅 ADMIN 可调用？"],
        ["3", "ProjectService.java", "事务中如何统计 unknown/skipped/invalid？"],
        ["4", "DecisionPolicy.java", "在线与离线迁移为何必须共享规则？"],
        ["5", "TaskRepository/DecisionRepository", "查询与唯一键如何支撑幂等？"],
        ["6", "ApiWorkflowIntegrationTests.java", "哪些状态变化由真实 HTTP 测试证明？"],
        ["7", "campus-deploy/*.ps1", "如何把环境、数据库、VM、日志和健康检查串起来？"],
        ["8", "capture_platform_screenshots.mjs", "如何不用 Playwright 也能稳定截图并释放任务？"],
    ], [900, 3700, 4760])

    doc.add_heading("24.7 你应该能独立复现的验收清单", level=2)
    add_bullets(doc, [
        "能从零构建 frontend/dist 和 standalone JAR，并解释静态资源进入 JAR 的位置。",
        "能用占位配置启动平台，不把任何真实密码、JWT 或数据库地址提交到 Git。",
        "能先 dry-run 再导入 30,183 条任务，并说明为何每批不超过 500。",
        "能重复迁移同一份历史决定，证明第二次 imported=0、skipped=N。",
        "能用两个独立账号完成领取、心跳、提交、释放和审计核对。",
        "能运行 Python 42 项测试、Java 6 项测试、Vue build 和 standalone package。",
        "能说清本轮验证证明了什么，以及尚未证明最大并发、灾备和公网安全。",
    ])
    add_callout(doc, "面试总结句", "我不仅实现了多人审核界面，还把已有桌面审核进度安全迁入在线状态机，在真实局域网中用独立账号验证了原子领取、租约和审计，并把部署与故障过程固化为脚本、测试和文档。", "info")

    chapter(doc, "25", "端到端案例：一张联合场景图如何变成可信训练数据", "用同一张 person + helmet + smoking 图片贯穿模型、规则、审核、并发、写回和评测。")
    add_picture(doc, ASSETS / "end-to-end-lifecycle.png", 6.6, "图 12  一张联合场景图片的完整数据生命周期", "从漏标源图片、单类 Teacher、几何分类、审核队列、原子领取、人工决定到派生数据集和重新评测的闭环")
    add_callout(doc, "案例设定", "图片 mine_000123.jpg 中真实存在 person、helmet、smoking；原标签只含 helmet。person Teacher 预测人体框 0.94，helmet Teacher 的框与原 GT 高度重叠，smoking Teacher 预测手口区域 0.78。", "info")

    doc.add_heading("25.1 阶段一：读取事实，但不急着改标签", level=2)
    add_paragraph(doc, "流水线先读取图片路径、原 YOLO 标签和数据集 names。它把 helmet GT 转成统一 Box 表示，并为每个 class/split 建立按图片索引。此时源目录没有任何写操作。模型版本、策略文件、classes、splits、imgsz 等组成 run signature，确保断点恢复不会把两次不同实验拼在一起。")
    add_table(doc, ["字段", "示例", "来源", "用途"], [
        ["split/image", "train/mine_000123.jpg", "数据集目录", "稳定定位图片与标签"],
        ["existing GT", "helmet [0.52,0.20,0.18,0.16]", "labels/train/*.txt", "判断预测是否已被同类标注覆盖"],
        ["model identity", "single_smoking_yolo26n_v1", "models JSON", "复现候选来自哪一个 Teacher"],
        ["policy version", "review_policy v2", "YAML", "解释当时使用的阈值与几何规则"],
        ["run signature", "SHA-256(normalized config)", "state.py", "阻止错误参数续跑"],
    ], [1800, 2600, 2300, 2660])

    doc.add_heading("25.2 阶段二：六个 Teacher 产生证据，不产生真值", level=2)
    add_paragraph(doc, "每次只加载一个单类别模型。person Teacher 扫到人体，helmet Teacher 扫到安全帽，smoking Teacher 扫到吸烟局部框。tractor、vest、slipper 没有预测。系统仍会统计 GT1_AUTO0、GT0_AUTO0 等空状态，因为‘模型没看到’也是评估 Teacher 覆盖能力的证据，但不会据此删除任何 GT。")
    add_table(doc, ["类别", "GT 数", "AUTO 数", "状态", "下一步"], [
        ["person", "0", "1", "GT0_AUTO1", "疑似漏标，进入几何/置信度判断"],
        ["helmet", "1", "1", "GT1_AUTO1", "检查是否同目标已标"],
        ["smoking", "0", "1", "GT0_AUTO1", "疑似小目标漏标"],
        ["vest/tractor/slipper", "0", "0", "GT0_AUTO0", "只统计，不生成候选"],
    ], [1500, 1100, 1300, 1800, 3660])
    add_callout(doc, "为什么枚举空状态", "如果只保存有预测的行，就无法区分‘图片没有该类’与‘图片有该类 GT 但 Teacher 漏检’。完整状态空间让模型能力和标签完整性分开分析。", "tip")

    doc.add_heading("25.3 阶段三：同类匹配与跨类关系必须分开", level=2)
    add_paragraph(doc, "helmet AUTO 与 helmet GT 先做同类 best_match。若 IoU/IoS/中心距离满足已标条件，关系为 GT_SAME_CONFIRMED，不生成新增框。person 和 smoking 没有同类 GT，因此不能因为它们与 helmet 或人体框重叠就直接删除；跨类别重叠只用于标记冲突或上下文，安全帽在人框内、吸烟框位于手口区域都是合理嵌套。")
    add_table(doc, ["候选", "同类 GT", "跨类关系", "几何结论", "推荐"], [
        ["helmet 0.91", "helmet GT", "位于 person 区域", "同目标已标", "不新增"],
        ["person 0.94", "无", "包含 helmet/smoking", "独立合法目标", "accept_add_or_reject"],
        ["smoking 0.78", "无", "位于 person 内且靠近 helmet", "跨类合理嵌套", "accept_add_or_reject"],
    ], [1700, 1900, 2000, 2100, 1660])
    add_callout(doc, "关键逻辑", "NMS 通常解决同一模型输出的重复框；本项目还要处理同类 GT/AUTO 重复、不同 Teacher 的跨类嵌套和原标签尺度歧义，不能把所有重叠都交给 NMS。", "warn")

    doc.add_heading("25.4 阶段四：候选进入审核包时包含什么", level=2)
    add_paragraph(doc, "候选 CSV 不是只有 class/conf/box。它还应携带 candidate_id、split、图片和标签相对路径、case_code、recommendation、同类最佳 GT、IoU/IoS/中心距离/面积比、跨类冲突、模型和策略版本、可视化文件。这样审核员能看图，系统能约束动作，事后能重放依据。")
    add_command(doc, """
    candidate_id: R-smoke-000123
    split: train
    image_name: mine_000123.jpg
    class_name: smoking
    confidence: 0.78
    case_code: GT0_AUTO1
    recommendation: accept_add_or_reject
    best_same_class_gt: null
    cross_class_context: person_contains_candidate
    visual_file: review_images/smoking/mine_000123__R-smoke-000123.jpg
    """, "这是概念化字段示例。真实字段以 review_queue.csv 为准；关键是候选身份、证据和允许动作必须一起流动。")
    add_paragraph(doc, "按图片聚合渲染会把 person、helmet、smoking 候选和原 GT 同时画在同一张审核图上。审核员看到的是联合场景，而不是三张割裂截图，因此能判断 smoking 小框是否被安全帽框视觉遮挡、是否为同一目标、是否符合标注规范。")

    doc.add_heading("25.5 阶段五：从 CSV 到数据库任务", level=2)
    add_numbers(doc, [
        "管理员创建项目，review_root 指向只读审核包。",
        "桥接脚本逐行验证 CSV，按不超过 500 条组成 HTTP 批次。",
        "Service 检查批内 candidate_id 重复，并查询数据库已有候选。",
        "review_tasks(project_id, candidate_id) 唯一键保证重复导入仍只有一条任务。",
        "导入结果显式返回 imported 和 skippedExisting，审计记录批次来源。",
    ])
    add_paragraph(doc, "此时 R-person-000123 与 R-smoke-000123 是两条 ReviewTask，但它们共享 image_name，前端可以把同图上下文展示给审核员。原 helmet 已确认，无需产生一个要求人工点击的新增任务。")

    doc.add_heading("25.6 阶段六：两个人同时审核为何不会串任务", level=2)
    add_paragraph(doc, "Reviewer A 与 Reviewer B 同时点击领取。数据库行锁让两次事务不能选中同一 PENDING 行；领取成功后，任务写入 claimed_by、lease_until 和 version。浏览器每 30 秒心跳续租，并显示剩余租约与网络状态。A 若关闭页面，租约过期后任务可回收；B 若拿旧 version 提交，服务端返回 409。")
    add_table(doc, ["时刻", "Reviewer A", "Reviewer B", "数据库不变量"], [
        ["T0", "请求 claim-next", "请求 claim-next", "同一任务最多一个事务持有写锁"],
        ["T1", "获得 person 候选", "等待/改选 smoking 候选", "candidate_id 不相同"],
        ["T2", "心跳续租", "提交 ACCEPT_ADD", "任务所有者和租约必须有效"],
        ["T3", "提交 ACCEPT_ADD", "刷新进度", "每个 task_id 最多一条 decision"],
    ], [1200, 2600, 2600, 2960])

    doc.add_heading("25.7 阶段七：人工接受后也不能直接 append", level=2)
    add_paragraph(doc, "审核完成后，review_apply.py 先确认没有空白或 UNCERTAIN 决策，再完整复制/硬链接源数据到新目录。对 ACCEPT_ADD，它重新读取目标标签、验证类别和框数值，并再次检查是否已出现同类重复；对 ACCEPT_REPLACE_GT，它验证被替换 GT 的引用和坐标仍与审核时一致。源 GT 若被别人修改，写回必须拒绝而不是猜测。")
    add_code(doc, "yolo_label_recovery/review_apply.py", 112, 195, "派生写回先检查审核是否完整，再处理 replace 引用与源 GT 变化，最后对新增框执行 duplicate_recheck。")
    add_table(doc, ["决策", "派生标签行为", "审计产物"], [
        ["ACCEPT_ADD person", "在新 labels/train 文件追加 person 框", "applied_or_replaced.csv"],
        ["ACCEPT_ADD smoking", "追加 smoking 小目标框", "applied_or_replaced.csv"],
        ["REJECT", "不改变标签", "保留 reviewer 和理由"],
        ["UNCERTAIN", "禁止生成最终数据集或转高级复核", "未解决清单/ESCALATED"],
    ], [2100, 4100, 3260])

    doc.add_heading("25.8 阶段八：新数据集是否更好必须靠固定测试集证明", level=2)
    add_paragraph(doc, "补标完成不等于模型一定提升。训练新版六类模型时应固定预训练起点、imgsz、优化器、数据划分和评测脚本，与上一版做可比实验。重点查看 smoking/slipper 等目标的分类别 Precision、Recall、mAP50、mAP50-95，以及 helmet+smoking、person+slipper 联合场景的定向样例。")
    add_table(doc, ["观察结果", "可能解释", "下一步"], [
        ["整体 mAP 升、smoking 不升", "补标主要改善 person/helmet，吸烟质量仍不足", "抽查 smoking 接受项和小目标尺度分布"],
        ["Recall 升、Precision 降", "补标增加覆盖但引入假阳性或阈值偏低", "分析误检并增加困难负样本"],
        ["val 升、test 不升", "划分泄漏、近重复或对验证集过拟合", "全局哈希/感知去重并固定独立 test"],
        ["指标变化小、联合场景变好", "总体平均掩盖关键场景收益", "保留场景级测试集和可视化证据"],
    ], [2100, 3700, 3560])
    add_callout(doc, "闭环结论", "系统价值不是把所有预测自动写入，而是把‘不完整监督’转化为可发现、可解释、可审核、可回滚、可评测的工程闭环。", "info")

    doc.add_heading("25.9 用这张图回答面试追问", level=2)
    add_bullets(doc, [
        "问算法：从 helmet 已标、person/smoking 未标，解释 GT/AUTO 状态与多几何量判断。",
        "问系统：从 candidate_id 解释 CSV、数据库唯一键、历史迁移和审计如何串联。",
        "问并发：从两账号领取解释行锁、租约、版本号各自保护哪个时刻。",
        "问安全：从审核图读取解释 JWT、项目成员、normalize/startsWith 和只读目录。",
        "问可靠性：从 OOM 和重复导入解释批次提交、run signature、幂等和回滚。",
        "问效果：从固定 test 和联合场景专项集解释为什么不能只报训练集曲线。",
    ])
    add_callout(doc, "自测标准", "不看手册，用 15 分钟把 mine_000123.jpg 从源标签讲到新版模型评测，并在每个阶段说出一个失败分支和一个验证证据。能做到，前后逻辑才真正连起来。", "tip")

    chapter(doc, "26", "用户反馈驱动的审核生产力迭代", "从真实审核速度、误操作恢复和前端可观测性出发，把功能堆叠改造成可持续使用的生产工作台。")
    add_picture(doc, ROOT / "docs" / "assets" / "platform-review-multibox.png", 6.6, "图 13  五候选按图聚合审核与租约状态", "真实运行平台：候选列表、完整图片、连接状态、受约束决策和快捷键同时可见")
    add_callout(doc, "本轮最重要的产品结论", "审核员每天处理的是大量重复判断。主流程每增加一次确认或一次手动跳转，都会按候选数量线性放大成本。安全设计必须保护数据，但不能无证据地牺牲吞吐。", "info")

    doc.add_heading("26.1 为什么‘先选择、再确认、再下一条’被回退", level=2)
    add_paragraph(doc, "第一版安全优化把一个决定拆成‘选择草稿、确认保存、手动下一框’。它降低了单次误点风险，却把每条候选从一次点击变成二至三次操作，并打断审核员连续判断的节奏。真实用户反馈表明，这个成本高于收益，因此主流程恢复为一键提交和自动前进。")
    add_table(doc, ["方案", "每条主要操作", "优点", "问题", "最终选择"], [
        ["二次确认", "选择 + 确认 + 下一条", "误点后提交前可撤销", "吞吐下降，注意力被 UI 打断", "不用于高频主流程"],
        ["一键审核", "点击决定", "速度快、认知连续", "误点需要后续纠错", "主流程采用"],
        ["最近审核纠错", "打开历史 + 改判", "不拖慢正常流，同时可恢复", "需要版本与权限控制", "作为安全补偿"],
    ], [1600, 1800, 2100, 2360, 1500])
    add_callout(doc, "设计原则", "不要把所有安全都做成阻塞式确认。高频路径优先防重复、防越权和可追溯；低频误操作通过最近历史、乐观版本和审计修订恢复。", "tip")

    doc.add_heading("26.2 一键保存与自动前进的状态流程", level=2)
    add_numbers(doc, [
        "审核员点击受 DecisionPolicy 约束的决定按钮，或按 1/2/3 快捷键。",
        "前端立即携带 task.id、expectedVersion、decision 和 comment 调用服务端。",
        "服务端校验任务所有者、租约、版本、允许动作和唯一决定，然后在事务内完成状态迁移。",
        "前端刷新本图候选与项目进度，优先领取同图下一条 PENDING 候选。",
        "本图没有待审框时，释放本地图像资源并调用 claim-next 自动进入下一图。",
        "历史决定改判不自动跳转，给审核员保留核对修订结果的上下文。",
    ])
    add_code(doc, "platform/frontend/src/App.vue", 160, 204, "submitDecision 区分主审核和历史改判；advanceAfterDecision 优先同图候选，再领取下一图。")
    add_callout(doc, "失败语义", "任何 API 校验失败都不会前进。busy 状态阻止重复点击；409 表示版本已变化，必须重新加载而不是覆盖。自动前进发生在服务端成功返回之后。", "warn")

    doc.add_heading("26.3 最近审核为什么不是普通历史列表", level=2)
    add_paragraph(doc, "最近审核接口按当前登录用户和当前项目同时过滤，limit 被限制在 1-50。它不是管理员全局审计的替代品，而是审核员自己的短期恢复入口。原审核人可改自己的决定，管理员可改全部决定；其他审核员只能查看同图已完成框，不能越权改判。")
    add_table(doc, ["保护层", "实现", "阻止的问题"], [
        ["身份范围", "reviewer_id = currentUser.id", "看到其他审核员的个人工作历史"],
        ["项目范围", "project_id + requireAccess", "跨项目读取候选和图片"],
        ["数量边界", "limit 夹在 1-50", "无界历史查询拖慢服务"],
        ["修订权限", "原审核人或 ADMIN", "任意成员修改别人的结论"],
        ["并发保护", "expectedVersion + @Version", "旧页面覆盖新决定"],
        ["审计", "TASK_DECISION_REVISED", "改判后无法追责"],
    ], [1900, 3600, 3860])
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TaskService.java", 146, 153, "recentDecisions 先做项目授权，再限制分页大小并按当前审核人查询。")

    doc.add_heading("26.4 租约、心跳和网络状态为什么必须可见", level=2)
    add_paragraph(doc, "以前心跳在后台静默运行，用户无法区分‘按钮没反应’、‘网络断开’和‘租约已经失效’。现在顶部显示浏览器在线状态，任务头显示剩余租约和续租状态；30 秒心跳在租约到期前持续延期。剩余不超过 120 秒或心跳异常时使用警示色。")
    add_table(doc, ["状态", "界面信号", "用户动作"], [
        ["连接正常", "绿色状态点 + 心跳正常", "继续审核"],
        ["正在续租", "显示续租中", "短暂等待，不重复点击"],
        ["浏览器离线", "红色网络离线", "检查校园网/网线，不继续提交"],
        ["租约不足 2 分钟", "倒计时警示色", "尽快提交或等待续租"],
        ["心跳失败", "心跳异常并显示错误", "重新领取，避免使用过期页面"],
    ], [1900, 3100, 4360])

    doc.add_heading("26.5 图片尺寸不是越大越好", level=2)
    add_paragraph(doc, "审核图贴满画布时，人物和框虽然更大，但四周上下文、工具栏和决策区同时被挤压。默认自适应因此改为画布的 88%，保留稳定留白；原始尺寸、50%-400% 缩放和全屏仍用于小目标细看。默认值服务于大多数判断，工具服务于例外。")

    doc.add_heading("26.6 为什么没有整体迁移 Element Plus", level=2)
    add_paragraph(doc, "Element Plus 擅长通用表格、分页、表单和弹窗，但本项目核心是图像画布、同图候选、键盘审核和状态提示组成的专用工作台。整体迁移会增加依赖、主题覆盖和关键交互回归成本，却不能自动解决吞吐、租约或纠错问题。因此保留自研主工作台；未来账号、项目、审计日志发展成复杂后台时，再按页面选择性引入组件。")
    add_table(doc, ["评估维度", "自研工作台", "整体迁移组件库", "决策"], [
        ["核心交互匹配", "为图像审核定制", "需要大量二次封装", "保留自研"],
        ["包体积", "当前生产 JS gzip 约 35 KB", "增加运行时与样式", "保持轻量"],
        ["视觉一致性", "已被真实用户验证", "需要重新主题化", "避免无收益回归"],
        ["管理后台扩展", "复杂表格成本会增长", "现成组件有优势", "未来按需采用"],
    ], [1800, 2500, 2800, 2260])

    doc.add_heading("26.7 查询索引如何从界面反推", level=2)
    add_paragraph(doc, "数据库索引不应只看字段频率，而应从真实查询路径出发。本图候选列表按 project_id、split、image_name 查找并按 id 排序；最近审核按 reviewer_id 和 decided_at 倒序。Flyway V2 为这两条路径增加复合索引，并已在 MySQL 8.4 的真实库上验证为 schema version 2。")
    add_table(doc, ["索引", "服务的交互", "没有索引的风险"], [
        ["(project_id, split, image_name, id)", "本图审核框左侧列表", "任务增长后按图查找反复扫描"],
        ["(reviewer_id, decided_at, id)", "我的最近审核", "每次打开历史都排序大量决定"],
    ], [3400, 2600, 3360])
    add_callout(doc, "面试表达", "这轮迭代不是‘又加了几个按钮’，而是用真实审核反馈重构吞吐与安全的边界：高频路径一键化，低频错误可恢复，后台状态可观察，查询路径有索引，设计取舍有 ADR。", "info")

    chapter(doc, "27", "把子项目讲回完整矿区系统", "从数据治理子系统出发，建立在线监控、业务处置和离线模型迭代的统一架构。")
    add_picture(doc, ASSETS / "mining-system-architecture.png", 6.65, "图 14  矿区智能安全监控系统双闭环架构", "在线检测与处置链路、离线数据与模型闭环")
    add_paragraph(doc, "标签恢复平台不是矿区业务的终点，而是模型生命周期中的数据治理子系统。在线链路负责从摄像头产生检测、事件、告警和工单；离线链路负责收集误报、漏报和联合场景困难样本，经过 Teacher 扫描、人工审核、派生数据集和固定测试集评估，再把合格模型部署回线上。")
    add_table(doc, ["层", "主要职责", "事实来源", "失败时的处理"], [
        ["设备与视频", "RTSP 拉流、抽帧、断流重连", "设备配置与通道状态", "单通道隔离重连，不拖垮全局"],
        ["模型与事件", "检测、跟踪、时序聚合", "frame_id、模型版本、原始框", "丢弃过期帧，保留关键证据"],
        ["告警与工单", "分级、派发、处置、关闭", "MySQL 事务状态", "幂等重试、状态机拒绝非法流转"],
        ["RAG 与 Agent", "查规程、生成建议、调用工具", "有效规程与工具执行记录", "证据不足拒答，高风险转人工"],
        ["数据与训练", "补标、审核、训练、评测", "只读源数据与派生版本", "可回滚，不污染原始标签"],
    ], [1500, 2600, 2600, 2660])
    add_callout(doc, "实施边界", "本仓库已有充分证据支撑的是 Multi-Teacher、多人审核平台、MySQL 协作控制和安全写回。视频业务、RAG、Agent 与边缘部署必须按真实完成度表述为已实现、原型或架构设计。", "warn")

    chapter(doc, "28", "Spring Boot 业务主干：设备、告警与工单", "用模块化单体、状态机和事务承载确定性业务规则。")
    add_paragraph(doc, "推荐先采用模块化单体：auth、device、detection、alert、workorder、knowledge、agent、audit 各自保持代码边界，但共享一次部署和本地事务。GPU 推理与视频解码因资源模型不同，可作为首批独立进程。")
    add_table(doc, ["模块", "核心实体", "关键不变量", "典型接口"], [
        ["device", "Area、Device", "RTSP 凭据不向前端明文暴露", "启停、状态、采样策略"],
        ["detection", "DetectionEvent", "每条结果携带设备、帧号、捕获时间和模型版本", "批量接收检测结果"],
        ["alert", "AlertRule、Alert", "单帧预测不能直接等于业务告警", "确认、驳回、升级"],
        ["workorder", "WorkOrder、WorkOrderLog", "状态流转受控，历史日志只追加", "接受、完成、升级、时间线"],
        ["audit", "AuditLog", "关键动作保留操作者、前后状态和时间", "按对象和用户查询"],
    ], [1600, 2200, 3100, 2460])
    add_command(doc, """
    PENDING -> CONFIRMED -> ASSIGNED -> PROCESSING -> RESOLVED -> CLOSED
                 |                                      |
                 +-> REJECTED                            +-> ESCALATED
    """, "状态不是由前端随意传入。Service 根据当前状态、角色、证据和版本号判断目标状态是否合法。")
    add_callout(doc, "事务边界", "创建工单时锁定告警、检查是否已有工单、写入工单、修改告警状态并追加日志应位于同一事务；数据库 UNIQUE(alert_id) 负责最终幂等兜底。", "tip")

    chapter(doc, "29", "20 路 RTSP、GPU 调度与端到端延迟", "用解耦、有界队列和最新帧优先把有限算力转化为稳定服务。")
    add_picture(doc, ASSETS / "rtsp-event-pipeline.png", 6.65, "图 15  多路视频推理与时序告警流水线", "最新帧、有界队列、动态 Batch、事件聚合和困难样本回流")
    add_paragraph(doc, "错误方案是按摄像头串行执行‘拉帧—推理—画框—返回’，或者每路加载一份模型。前者产生头部阻塞，后者造成 CUDA Context 和权重副本膨胀。推荐每路独立拉流，只保留一个最新待处理帧，经共享有界队列进入单 GPU 模型实例。")
    add_table(doc, ["机制", "解决的问题", "关键选择", "监控指标"], [
        ["最新帧槽", "历史帧积压导致假实时", "新帧覆盖未处理旧帧", "丢帧率、帧龄"],
        ["有界队列", "内存无限增长", "满载时丢旧帧或低优先级帧", "队列长度、等待 P95"],
        ["动态 Batch", "吞吐与等待冲突", "数量到阈值或等待到时立即执行", "平均 Batch、推理 P95"],
        ["公平调度", "单路摄像头占满队列", "每通道最多一个待处理帧", "各路实际推理 FPS"],
        ["frame_id 对齐", "旧框画到新视频帧", "结果过期则丢弃或同帧后端画框", "端到端 P95"],
    ], [1600, 2400, 3100, 2260])
    add_formula(doc, "容量估算", "需求 FPS = 摄像头路数 × 每路检测 FPS", "20 路每路 2 FPS 需要 40 FPS 推理能力。系统不能长期运行在 100% 利用率，应为疑似事件升频和抖动保留余量。")
    add_callout(doc, "延迟口径", "模型推理延迟、捕获到推理完成、捕获到前端展示是三个指标。单帧低于 80 ms 不能直接证明 20 路端到端没有延迟。", "warn")

    chapter(doc, "30", "吸烟与拖鞋：从小目标检测到可靠事件", "用空间关联、时序窗口、迟滞和去重抵抗单帧波动。")
    add_paragraph(doc, "Detection、Track、Event、Alert 必须分开：单帧框是模型输出，轨迹关联同一人员，事件表示一段持续行为，告警是满足业务规则后的持久记录。同一个人连续吸烟 30 秒可以产生几十个框，但只应形成一个事件和一条告警。")
    add_table(doc, ["阶段", "吸烟示例", "拖鞋示例", "目的"], [
        ["空间关联", "框位于人员头肩/手口区域", "框位于人员底部/脚部区域", "过滤人体结构不合理误检"],
        ["轨迹关联", "绑定 person track_id", "绑定 person track_id", "避免把不同人的命中合并"],
        ["时间窗口", "20 秒内至少 3 次", "多个采样时刻重复出现", "容忍间歇漏检"],
        ["迟滞", "严格进入、宽松维持、长时间无命中退出", "同理", "防止事件反复开关"],
        ["去重", "设备+类别+轨迹+时间桶", "同理", "阻止重复告警"],
    ], [1500, 2900, 2800, 2160])
    add_formula(doc, "事件评分示例", "S = 0.45H + 0.30C + 0.15T + 0.10G", "H 为窗口命中比例，C 为平均置信度，T 为时间覆盖率，G 为空间关联质量。它是业务评分而非模型概率，必须由现场数据校准。")
    add_callout(doc, "联合场景排查", "安全帽出现后吸烟消失，不应先归因于类别框重叠。要依次检查输入是否已画框、是否 class-agnostic NMS、640 下目标是否过小、阈值是否过高，以及训练集中安全帽+吸烟联合标注是否完整。", "warn")

    chapter(doc, "31", "RAG 安全知识库：证据先于答案", "把规程版本、混合检索、重排序和风险控制串成可信问答链路。")
    add_paragraph(doc, "知识入库不是把 PDF 扔进向量库。系统先做文件哈希去重、OCR、页眉页脚清洗、标题层级恢复和条款级语义切片，并保留文档标题、版本、生效状态、章节、条款号和页码。旧版本标记为 INACTIVE，但历史工单仍能访问当时引用的版本。")
    add_numbers(doc, [
        "向量检索召回语义相近的规程，BM25 召回精确编号、设备名和专业词。",
        "使用 RRF 融合两个排序列表，避免直接比较不同量纲的分数。",
        "Cross-Encoder 对少量候选重排序，最终只把最相关证据交给大模型。",
        "回答必须返回规程、章节、页码和原文片段，前端可跳回原文核验。",
        "高风险问题证据不足时拒绝生成具体操作，转人工确认。",
    ])
    add_formula(doc, "Reciprocal Rank Fusion", "RRF(d) = Σ 1 / (k + rank_i(d))", "同一知识块在多个召回列表中排名越靠前，融合分数越高；它不要求 BM25 和向量相似度处于同一尺度。")
    add_table(doc, ["风险", "证据要求", "系统行为"], [
        ["LOW", "一般制度材料", "可回答并标注来源"],
        ["MEDIUM", "至少一条有效规程", "生成建议，必要时人工确认"],
        ["HIGH", "有效规程原文且与场景匹配", "证据不足拒答并升级人工"],
    ], [1500, 3600, 4260])

    chapter(doc, "32", "Function Calling Agent：大模型负责理解，代码负责安全", "用工具白名单、状态机、幂等和人工审批限制大模型权力。")
    add_paragraph(doc, "Agent 不直接连接数据库，也不生成 SQL 执行。模型只能从固定工具中选择调用，并提交符合 JSON Schema 的参数；Spring Boot 仍负责认证、授权、参数校验、事务和审计。Prompt 不是安全边界，真正的边界在工具网关和 Service。")
    add_table(doc, ["工具", "是否有副作用", "服务端必须校验", "失败策略"], [
        ["query_safety_knowledge", "否", "项目范围、文档状态、风险过滤", "可有限重试"],
        ["get_device_info", "否", "设备访问权和脱敏字段", "可有限重试"],
        ["create_work_order", "是", "告警状态、唯一性、权限、负责人", "按幂等键查询结果"],
        ["escalate_alert", "是", "当前风险、操作者权限和状态流转", "失败转人工"],
        ["request_human_approval", "是", "审批人范围与有效期", "等待或终止"],
    ], [2100, 1500, 3700, 2060])
    add_command(doc, """
    RECEIVED -> COLLECTING -> RETRIEVING -> ASSESSING
       -> WAITING_APPROVAL -> DISPATCHING -> MONITORING -> COMPLETED
    """, "状态机限制当前阶段允许调用的工具；高风险动作不能跳过 WAITING_APPROVAL。")
    add_paragraph(doc, "每个有副作用的工具调用携带 idempotencyKey = agentRunId + stepId + toolName。网络超时后先查询执行状态，不能盲目再次创建工单。RAG 文档和用户输入都视为不可信内容，文档中的‘忽略系统规则’不得改变工具白名单。")

    chapter(doc, "33", "MySQL：把并发、状态和审计变成约束", "理解表关系、事务隔离、行锁、乐观版本、索引和恢复。")
    add_paragraph(doc, "MySQL 保存的是需要强一致的业务事实：用户与项目成员、审核任务与决定、设备、告警、工单、Agent 调用和审计记录。图片与视频证据更适合文件系统或对象存储，数据库保存路径、哈希、时间和关联对象。")
    add_table(doc, ["查询路径", "推荐索引", "支撑的界面/业务"], [
        ["领取审核任务", "(project_id, state, lease_until, id)", "原子 claim-next"],
        ["同图候选", "(project_id, split, image_name, id)", "左侧本图所有框"],
        ["最近审核", "(reviewer_id, decided_at, id)", "误操作修订"],
        ["设备告警", "(device_id, status, created_at)", "告警列表与时间线"],
        ["事件幂等", "UNIQUE(device_id, event_key)", "阻止重复告警"],
    ], [2400, 3900, 3060])
    add_paragraph(doc, "行锁只存在于领取任务的短事务中，租约跨请求保存业务所有权，@Version 在提交时阻止旧页面覆盖新状态。三者分别保护毫秒级竞争、分钟级占用和提交时并发覆盖，不能相互替代。")
    add_callout(doc, "数据库面试重点", "会用 EXPLAIN 观察 type、key、rows、filtered 和 Extra；知道深分页应改用 Keyset Pagination；事务内不做模型推理或远程调用；死锁通过统一锁顺序、缩短事务和有限重试处理。", "tip")

    chapter(doc, "34", "Redis、Kafka 与微服务：按瓶颈演进，不堆技术", "说明何时需要中间件，以及如何避免缓存、消息和数据库之间的不一致。")
    add_table(doc, ["能力", "适合保存/处理", "不应承担", "引入信号"], [
        ["Redis", "设备心跳、滑动窗口、限流、热点缓存", "审核最终决定等核心事实", "高频短状态压垮数据库"],
        ["Kafka", "告警通知、统计、证据归档、困难样本回流", "用户必须立即获知结果的领取请求", "多下游异步订阅和明显积压"],
        ["微服务", "独立扩缩容的视频、推理、知识服务", "尚未稳定的小模块", "团队边界、发布频率和资源模型分化"],
        ["Kubernetes", "多服务编排、滚动发布、弹性和自愈", "单机演示的复杂度装饰", "已有容器化、监控和多节点需求"],
    ], [1500, 3100, 3000, 1760])
    add_paragraph(doc, "数据库与 Kafka 不能直接双写。推荐 Transactional Outbox：同一 MySQL 事务写业务表和 outbox_event，后台发布器再把未发布事件发送到 Kafka。消费者以 event_id 幂等处理，并将业务更新与 consumed_event 记录放在同一事务。")
    add_command(doc, """
    阶段 1  模块化单体 + MySQL
    阶段 2  独立 GPU 推理服务 + 对象存储
    阶段 3  Redis 承载心跳、窗口和热点状态
    阶段 4  Kafka 解耦告警、通知和样本回流
    阶段 5  按扩容与团队边界拆分微服务
    """, "这是一条演进路线，不是当前仓库已使用技术的清单。")

    chapter(doc, "35", "大厂面试表达：三分钟讲清、三十分钟扛住", "用业务问题、技术取舍、验证证据和诚实边界组织项目故事。")
    add_callout(doc, "30 秒版本", "我负责矿区视觉模型的数据治理和多人审核平台。针对多来源六类 YOLO 数据中的联合场景漏标，我训练六个单类别 Teacher 全量扫描，用 IoU、IoS、中心距离和面积比例识别疑似漏标，再通过 Spring Boot、MySQL、Vue 多人平台安全审核与回写，形成可追溯的数据闭环。", "info")
    add_paragraph(doc, "三分钟版本按六段组织：业务问题；为什么漏标会产生错误负监督；Multi-Teacher 与流式资源控制；多人协作与数据库一致性；固定测试集和联合场景验证；已落地能力与下一步边界。不要按 Python、Java、Vue 的技术清单顺序介绍。")
    add_table(doc, ["追问", "回答主线", "不要踩的坑"], [
        ["最难的是什么", "先区分模型、数据和部署链路，再定位联合场景漏标", "只说调了学习率和 Batch"],
        ["为何单类 Teacher", "减少类别竞争、按类校准、专注发现漏标", "说单类模型天然正确"],
        ["为何高置信仍审核", "域偏移与确认偏差，伪标签会放大错误", "把置信度当真实概率"],
        ["指标没提升怎么办", "查划分、标签质量、分类别指标和专项场景", "继续盲目加轮次"],
        ["为何不用 Redis/Kafka", "当前 MySQL 单体满足一致性，按瓶颈演进", "冒充已经使用"],
        ["如何证明可用", "固定 test、联合场景、双账号并发和可复核产物", "用训练损失或截图代替评测"],
    ], [2300, 4200, 2860])
    add_bullets(doc, [
        "可以确认：30,183 个候选任务导入，4,465 条历史决定迁移，双账号校园网协作已验证。",
        "不能夸大：当前不应称为高并发生产平台；Redis、Kafka、Kubernetes 是演进方案。",
        "口径分离：模型单帧延迟不等于 20 路端到端延迟；mAP 不等于事件级告警准确率。",
        "责任边界：明确自己负责数据治理、模型训练评测和审核平台，其他模块按参与程度陈述。",
    ])
    add_callout(doc, "STAR 收束", "Situation：多源六类数据联合场景漏标。Task：降低人工成本且不污染原标签。Action：六 Teacher、几何分流、流式推理、多人审核、事务与审计。Result：形成可运行、可追溯、可继续评测的数据闭环；模型收益仍以固定测试集和现场专项集为准。", "tip")

    chapter(doc, "36", "从头到尾讲解 GitHub 工程作品集", "按 README 的真实阅读顺序理解每一段文字、每一张图、每一种证据及其面试用途。")
    add_paragraph(doc, "大厂面试官通常不会从仓库第一行读到最后一行。首页必须先回答：解决什么真实问题、最难的工程矛盾是什么、有哪些真实运行证据、如何在没有私有数据和 GPU 的情况下复现、哪些能力仍是下一阶段设计。当前仓库使用英文作为 GitHub 默认首页，降低国际面试官的阅读门槛；中文长版、代码导读和本手册则负责承载更深入的原理、实验与面试复盘。")
    add_table(doc, ["停留时间", "首页应提供的证据", "面试官形成的判断"], [
        ["30 秒", "一句话价值、技术栈、真实审核界面", "不是孤立脚本，而是完整工程作品"],
        ["5 分钟", "完整链路、预生成报告、生产规模数字", "既有算法，也有系统与质量治理"],
        ["15 分钟", "无 GPU 演示、预期输出和失败验收", "第三方能够复现关键机制"],
        ["30 分钟", "核心代码导读、事务时序和数据不变量", "能够接受后端、算法和系统设计深挖"],
    ], [1500, 3900, 3960])

    doc.add_heading("36.1 首页开场：名称、定位、徽章与主视觉", level=2)
    add_paragraph(doc, "README 顶部先给出项目名 YOLO Label Recovery，再用一句英文明确产品边界：这是一个安全、可审计、内存友好的 Multi-Teacher 漏标恢复与人工审核平台，而不是一个只会调用 YOLO predict 的脚本。Python、Spring Boot、Vue、CI、Release 和 MIT 徽章分别对应算法流水线、多人后端、审核前端、持续集成、可交付版本和开源许可。徽章的意义是快速建立技术栈和工程成熟度预期，不是装饰。")
    add_picture(doc, ROOT / "docs" / "assets" / "platform-review-multibox.png", 6.65, "图 16  GitHub 首页主视觉：真实五候选多人审核工作台", "真实地下场景中按图聚合三个人员、一个背心和一个拖拉机候选；同时展示候选列表、租约、置信度、决策动作、快捷键与释放任务入口")
    add_callout(doc, "这张主图应该怎么看", "先看左侧：同一图片的五个候选被聚合而不是拆成五个页面；再看中间：审核员始终保留完整场景语义；最后看右侧：候选证据、租约、三种受约束决策和快捷键全部可见。它证明产品形态真实存在，但不能用来证明模型精度或高并发能力。", "info")
    add_callout(doc, "30 秒讲法", "这是我们为矿区六类检测数据构建的漏标恢复平台。模型只产生候选，规则解释 GT 与 AUTO 的关系，人工最终授权；Web 平台再用账号、原子领单、租约、版本和审计保证多人审核不会互相覆盖。", "tip")

    doc.add_heading("36.2 Pre-generated showcase：为什么没有 GPU 也能验收", level=2)
    add_paragraph(doc, "主视觉之后进入预生成展示区。这里的分析报告来自仓库内置的合成 fixture，目的是让任何面试官无需私有图片、模型权重和 GPU，也能复现输入契约、状态分类、阈值策略、资源恢复和安全写回。真实审核界面则使用经许可的生产审核样例。两者共同证明工具行为，但都不能冒充新模型的固定测试集精度。")

    doc.add_heading("36.2.1 Model-free dataset audit：先证明数据能安全进入流水线", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "audit-preview.png", 6.35, "图 17  无模型数据集审计报告", "合成数据主动注入非法类别、孤立标签和 train/val 跨划分重复，报告必须返回 FAIL 才算验收成功")
    add_paragraph(doc, "这张图回答的是‘输入数据是否可信’，而不是‘模型是否准确’。审计会检查图片/标签配对、类别 ID、坐标合法性、损坏文件、空标签、精确哈希重复和跨划分泄漏。公开 fixture 的预期结果是 FAIL、两个严重问题、一个警告和一组跨划分重复。故意失败代表审计器抓住了注入缺陷；如果错误返回 PASS，才说明实现有问题。")
    add_callout(doc, "面试追问", "为什么训练前必须做哈希和跨划分检查？因为相同图片同时进入 train 与 val 会造成数据泄漏，让验证指标虚高；错误标签则会把数据问题伪装成模型或超参数问题。", "info")

    doc.add_heading("36.2.2 Multi-Teacher recovery report：证明全量扫描可执行、可恢复", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "report-preview.png", 6.35, "图 18  Multi-Teacher 漏标恢复质量报告", "集中展示图片-模型扫描次数、分类别候选、AUTO/REVIEW、稳定 Batch、OOM 重试和源标签只读状态")
    add_paragraph(doc, "报告中的 image-model scans 不是图片数，而是图片数量乘以 Teacher 数量。它同时记录每个模型最终稳定 Batch 和模拟 OOM 重试，说明实现采用单模型串行加载、流式图片读取和批次级提交，而不是把六个模型和全部图片一次塞进显存或内存。Source labels modified = No 是关键不变量：扫描只生成证据和派生结果，不覆盖原始标签。")

    doc.add_heading("36.2.3 Audited threshold calibration：阈值来自审核证据", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "calibration-preview.png", 6.35, "图 19  分类别阈值校准报告", "使用六类共 2,400 条已审核候选估计 AUTO 精度下限和 REVIEW 召回覆盖，不使用一个全局阈值")
    add_paragraph(doc, "横向比较六个类别可以看到阈值并不相同：tractor 的 AUTO 阈值可以较低，而 smoking 需要更保守。AUTO 不只看样本精度点估计，而要求 95% Wilson 置信下限达到目标精度；REVIEW 阈值则优先保留足够多的审核正样本。这样做把‘我觉得 0.7 可以’升级为由历史审核数据支持、可复算的策略。")
    add_callout(doc, "不能夸大", "校准报告只对当前审核样本分布有效。数据域、模型版本或类别定义变化后必须重新抽样校准，置信度也不等于真实概率。", "warn")

    doc.add_heading("36.2.4 Cross-Teacher consensus：高置信还要独立证据", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "consensus-preview.png", 6.35, "图 20  跨 Teacher 一致性门控报告", "主 Teacher 的 AUTO 候选需要独立验证 Teacher 在空间上提供一对一支持，否则降级到 REVIEW")
    add_paragraph(doc, "主模型和验证模型分别产生候选后，系统按类别与空间关系做一对一匹配。公开 fixture 中 72 个主 AUTO 只有 48 个获得支持并保留 AUTO，其余 24 个降级人工复核。这个阶段处理已经落盘的候选证据，不会同时加载两套模型，因此不增加推理显存峰值。它降低单模型确认偏差，但不能保证两个模型不会犯同一种错误。")

    doc.add_heading("36.2.5 Perceptual near-duplicate groups：减少重复审核与数据泄漏", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "near-duplicates-preview.png", 6.35, "图 21  感知近重复分组报告", "聚合缩放、JPEG 重压缩和亮度变化版本，同时避免把纯黑、纯白等低纹理图片错误合并")
    add_paragraph(doc, "全局文件哈希只能发现字节完全相同的图片，感知哈希则用于发现视觉内容近似的版本。报告把七张图片归为三组，只需优先审核三个代表图，并标记跨 train/val/test 的近重复组。它既降低连续帧造成的审核浪费，也为重新划分数据提供泄漏证据。最终是否批量继承决定仍需检查组内一致性，不能仅凭相似就自动复制标签。")

    doc.add_heading("36.2.6 Diversity-aware active review：有限人力先看什么", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "prioritization-preview.png", 6.35, "图 22  多样性感知主动审核排序报告", "在有限审核预算下同时考虑类别稀缺、候选风险和感知多样性，避免 person 淹没 smoking 等小类")
    add_paragraph(doc, "候选达到几万条时，CSV 原始顺序没有业务意义。优先级模块先保证六个类别都能进入队列前部，再利用动态稀缺度照顾小类，并用感知距离抑制重复场景。公开 fixture 中预算 12 覆盖全部六类，前六个位置每类各一个。它优化的是审核资源分配，不改变候选真假，也不代替人工决定。")

    doc.add_heading("36.2.7 Exhaustive GT/AUTO gate：决定候选为什么需要人看", level=3)
    add_picture(doc, ROOT / "docs" / "assets" / "review-gate-preview.jpg", 5.9, "图 23  GT/AUTO 多几何关系审核图", "联合 IoU、IoS、归一化中心距离和面积比区分已标、包含、同目标尺度歧义、独立漏标、模型重复和跨类别冲突")
    add_paragraph(doc, "这张图是候选进入审核平台前最关键的解释层。GT0_AUTO0、GT1_AUTO0、GT0_AUTO1、GT1_AUTO1 枚举图片/类别的四种基础状态；IoU 衡量整体重叠，IoS 识别小框被大框包含，归一化中心距离判断中心是否接近，面积比识别同中心但尺度悬殊。再结合模型内部重复与跨类别冲突，系统生成关系、推荐动作和允许动作。高置信只影响证据强度，不能越过人工授权直接写标签。")
    add_table(doc, ["关系", "图中含义", "处理原则"], [
        ["确认已标 / 包含式已标", "AUTO 与同类 GT 对应同一目标", "忽略新增，避免重复框"],
        ["同目标歧义 / 包含式歧义", "中心接近但框尺度或边界差异大", "允许替换、保留或拒绝，由人工判断"],
        ["独立漏标", "AUTO 与现有同类 GT 空间上是另一个目标", "进入补标签审核"],
        ["模型内部重复", "多个 AUTO 指向同一目标", "只保留一个候选或人工去重"],
        ["跨类别冲突", "与其他类别高度重叠且不符合合理嵌套", "升级复核，不自动写入"],
    ], [2100, 4100, 3160])

    doc.add_heading("36.3 两代审核工具：从离线交付到多人协作", level=2)
    add_paragraph(doc, "README 接下来展示两代审核产品。旧版 Tk 桌面工具没有被淘汰，它仍适合单人、离线、拷贝即用的交付；Web 平台解决的是多人登录、原子领单、租约恢复、项目隔离、版本冲突与审计追踪。保留两代截图能够说明需求如何推动架构演进。")
    add_picture(doc, ROOT / "docs" / "assets" / "grouped-review-preview.png", 6.25, "图 24  旧版 Tk 按图片聚合审核器", "左侧列出本图全部候选，中心显示真实审核图，按钮受决策策略约束；支持中英文、快捷键、JSONL 日志、原子 CSV 检查点和崩溃续审")
    add_picture(doc, ROOT / "docs" / "assets" / "grouped-review-joint-scene.jpg", 5.9, "图 25  桌面审核器真实联合场景", "同一人物相关的多个类别候选必须放回完整场景判断，不能逐框脱离上下文审核")
    add_picture(doc, ROOT / "docs" / "assets" / "grouped-review-multibox-scene.jpg", 5.9, "图 26  桌面审核器真实多框矿区场景", "一张矿区图片中存在多个候选，按图聚合减少重复加载并帮助识别跨类别关系")
    add_paragraph(doc, "桌面端已经解决了按图聚合、动作约束、自动保存与恢复，但多人共享 CSV 会出现重复领取、最后写入覆盖、账号责任不清和跨项目数据暴露。因此下一张截图不是单纯把 Tk 改成网页，而是把协作一致性提升为服务端事务。")
    add_picture(doc, ROOT / "docs" / "assets" / "platform-login.png", 6.25, "图 27  Web 平台真实登录入口", "JWT 认证建立用户身份；ADMIN、REVIEWER、AUDITOR 角色决定全局能力，项目成员关系进一步限制可访问的数据范围")
    add_picture(doc, ROOT / "docs" / "assets" / "platform-dashboard.png", 6.25, "图 28  Web 平台真实项目看板", "展示生产规模任务总量、待审核、占用、疑难升级和完成率；这些数字证明导入与协作流程规模，不代表模型精度")
    add_picture(doc, ROOT / "docs" / "assets" / "platform-admin.png", 6.25, "图 29  Web 平台账号与项目成员管理", "管理员创建账号、分配角色并授权项目；只有创建账号而不加入项目，审核员仍无法领取该项目任务")
    add_callout(doc, "Web 并发主线", "claim-next 在数据库事务中使用悲观写锁领取任务；浏览器周期性续租，断开后租约过期可重新领取；提交决定携带乐观版本，旧页面返回 409 而不是覆盖新状态；每次领取、决定、纠错和释放都写审计事件。", "info")
    add_callout(doc, "真实验证边界", "平台已用两个独立账号在可信校园网完成同时审核，并导入 30,183 条任务、迁移 4,465 条历史决定。这证明任务分配、迁移和审计闭环可运行，但不能称为互联网级高并发压测。", "warn")

    doc.add_heading("36.4 Production-scale validation：真实数字应该如何解释", level=2)
    add_paragraph(doc, "README 的生产规模汇总使用脱敏 SVG，核心数字是 29,071 张私有六类图片、六个单类别 Teacher、174,426 次 image-model passes、99,696 条候选证据、30,183 条人工审核任务和 0 次渲染失败。它证明流水线能够完成真实规模扫描与审核包生成，且源标签未被修改；它不公开私有图片、权重和机器路径，也不声明候选都正确。")
    add_table(doc, ["生产证据", "能证明什么", "不能证明什么"], [
        ["174,426 次 image-model passes", "六 Teacher 完成全量串行扫描", "六模型同时驻留显存"],
        ["99,696 条候选证据", "关系分类与审计流有真实规模输入", "99,696 个都应补标"],
        ["30,183 条审核任务", "审核门控和平台导入可工作", "平台已经过高并发压测"],
        ["0 次渲染失败", "审核可视化产物完整", "模型没有误检和漏检"],
        ["源标签未修改", "实验可回滚、证据链可追溯", "派生标签无需人工验收"],
    ], [2500, 3900, 2860])

    doc.add_heading("36.5 Why、Architecture、Properties：从业务问题走到工程约束", level=2)
    add_paragraph(doc, "Why this project exists 用 person + helmet + smoking、person + slipper 等联合场景解释错误负监督：图片里真实存在目标却没有标签时，训练器会把它当背景。随后 Mermaid 流程图从只读数据集开始，依次经过审计、单 Teacher 加载、流式 FP16 推理、同类 GT 匹配、分类别置信路由、可选共识、GT/AUTO 审核门控、人工决定和派生数据集。这个顺序对应真实代码和数据状态，而不是理想化架构图。")
    add_bullets(doc, [
        "Immutable source：原标签目录永远不是输出目录，失败实验可以删除派生目录重来。",
        "One Teacher at a time：显存只保留当前模型与当前 Batch，六个模型不会同时常驻。",
        "Streaming and bounded memory：图片按批读取，候选按行落盘，复核图数量有上限。",
        "Adaptive Batch：OOM 时丢弃未提交批次、Batch 减半重试，成功后才推进检查点。",
        "Idempotent resume：run signature 防止错误参数续跑，candidate_id 和标签写入避免重复。",
        "Human authorization：AUTO 是高置信候选分段，不代表绕过审核直接修改源数据。",
        "Auditable collaboration：服务端保存领取人、租约、决定、版本和审计事件。",
    ])

    doc.add_heading("36.6 Quick start、输出目录和资源模型：面试官如何亲自复现", level=2)
    add_paragraph(doc, "README 的 Quick start 先给无 GPU 演示，再给可选 GPU 扫描。正确展示顺序不是一上来安装 CUDA，而是先运行 synthetic audit 和 review fixture，验证输入契约、预期失败、候选关系、人工动作和安全写回；只有讨论真实 Teacher 推理时才安装 inference 依赖。")
    add_command(doc, r"""
    python examples\create_synthetic_dataset.py --output .demo-dataset
    yolo-label-recovery audit .demo-dataset --output-dir .demo-audit --hash-images --check-images
    python examples\create_review_fixture.py --output-dir .demo-review-fixture
    yolo-label-recovery review-build .demo-review-fixture\dataset .demo-review-fixture\candidates.csv `
      --output-dir .demo-review-result --render --redact-paths
    yolo-label-recovery review-ui .demo-review-result
    """, "无 GPU 也能演示数据审计、GT/AUTO 关系、按图审核和失败验收；不要把预生成 fixture 当成私有模型精度测试。")
    add_paragraph(doc, "输出目录本身就是审计协议：candidates_all.csv 保存完整候选流，candidates_auto.csv 与 candidates_review.csv 保存路由结果，summary 解释统计，state.json 保存原子断点，manifest.json 记录参数、依赖、CUDA 与 GPU，report.html 提供可视化，trainable_dataset 只有显式物化时才生成。发布 GitHub 报告时必须 redact paths，避免泄露本地数据和权重路径。")
    add_table(doc, ["资源", "运行时实际保留", "避免的问题"], [
        ["GPU 显存", "当前 Teacher、当前 Batch 激活与预测张量", "六模型同时加载导致 OOM"],
        ["CPU 内存", "图片路径、当前解码批次、同类 GT 缓存、有限复核样本", "全量图片与可视化长期驻留"],
        ["磁盘", "流式 CSV、状态文件、派生标签和审核图", "内存无限增长且中断丢失进度"],
    ], [1700, 4700, 2960])

    doc.add_heading("36.7 文档导航、项目状态与 README 收尾", level=2)
    add_paragraph(doc, "README 最后把读者引向专题文档，并明确该仓库是经过清理的工程作品而不是公开模型 benchmark。真实图片、权重、日志和机器路径不会上传；公开 fixture 用来验证机制，真实固定测试集才用于声明模型收益。这样既能展示真实工程能力，又不泄露公司或项目数据。")
    add_table(doc, ["文档", "解决的问题", "面试用途"], [
        ["README.md", "项目是什么、为什么值得继续看", "英文开场与作品展示"],
        ["README.zh-CN.md / docs/README.zh-CN.md", "中文完整链路与不同读者入口", "快速切换讲解深度"],
        ["REPRODUCIBLE_DEMO.zh-CN.md", "如何在无 GPU 环境复现", "现场演示与验收"],
        ["CODE_WALKTHROUGH.zh-CN.md", "候选如何穿过 Python、Java、MySQL 和 Vue", "代码级追问"],
        ["COLLABORATION_PLATFORM.md", "登录、领单、租约、版本和审计如何工作", "后端与并发追问"],
        ["PROJECT_EVIDENCE.zh-CN.md", "哪些是实现、验证、原型和规划", "可信边界"],
        ["MINING_SYSTEM_DESIGN.zh-CN.md", "如何接入完整矿区监控、RAG 和 Agent", "系统设计扩展"],
        ["INTERVIEW_QA.zh-CN.md", "高频问题如何回答", "模拟面试"],
    ], [2800, 3800, 2560])
    add_callout(doc, "从头到尾的收束句", "这个 GitHub 不是在展示一组漂亮截图，而是在展示一条证据链：数据先被审计，Teacher 只提出候选，统计策略控制风险，多几何关系解释冲突，人工平台授权决定，派生数据安全写回，固定测试集验证收益；每个环节都有代码、报告、测试或真实运行截图支撑。", "tip")
    add_callout(doc, "作品集原则", "截图证明产品真实存在，测试证明代码行为，数据库约束证明并发边界，生产规模数字证明流程承载量，固定测试集证明模型收益。五类证据不能互相替代。", "warn")

    chapter(doc, "37", "沿候选生命线读懂核心代码", "不按目录背类名，而是追踪一条候选从 CLI、Teacher、审核到派生标签的完整状态变化。")
    add_logic_bridge(doc, "第 36 章解决别人如何看懂仓库", "自己必须能把首页中的每个结论追到代码", "用真实函数解释输入、状态、不变量和失败分支")
    add_paragraph(doc, "第一条阅读主线从 CLI 开始。一级命令只负责分发，run 才延迟加载 Ultralytics，因此数据审计、阈值校准和审核可以在没有 PyTorch 的环境执行。这是运行时依赖隔离，而不是单纯的代码风格。")
    add_code(doc, "yolo_label_recovery/cli.py", 40, 83, "观察 _load_pipeline 如何把可选 GPU 依赖限制在 run 命令，以及 main 如何把一级命令分发到独立模块。")
    add_paragraph(doc, "第二条主线进入 Teacher 扫描。参数、数据、模型和输出路径先验证，再建立 run signature；每个类别只加载一个模型和同类 GT 缓存。当前 batch 只有在候选、派生标签和断点都能安全提交后才算完成。")
    add_code(doc, "autolabel_with_single_class_models.py", 575, 650, "这一段完成数据契约、类别映射、模型权重和输出目录验证，并只扫描一次图片清单供六个 Teacher 复用。")
    add_code(doc, "yolo_label_recovery/state.py", 14, 55, "run signature 防止用不同数据、模型或参数继续旧实验；state.json 使用原子替换保存已提交游标。")
    add_paragraph(doc, "第三条主线是候选决策。模型输出不是标签，classify_candidate 联合同类和跨类 GT、模型内部重复和分类别置信度策略，生成关系、推荐动作和允许动作。")
    add_code(doc, "yolo_label_recovery/review_decision.py", 100, 174, "高置信度只影响置信分段；真正允许什么审核动作由 GT/AUTO 状态、几何关系和冲突共同决定。")
    add_paragraph(doc, "第四条主线是安全写回。人工决定仍不是直接 append；review-apply 再验证决策完整性、框合法性、源 GT 漂移和重复关系，并物化到新的数据集目录。")
    add_code(doc, "yolo_label_recovery/review_apply.py", 78, 146, "写回前建立稳定候选索引，拒绝未决项和非法动作，并保持源数据集只读。")
    add_table(doc, ["对象", "产生位置", "核心字段", "不变量"], [
        ["Candidate", "Teacher 扫描", "split、image、class、box、confidence", "candidate_id 稳定"],
        ["Review row", "review-build", "几何量、关系、允许动作", "证据可追溯"],
        ["Decision", "桌面/Web 审核", "动作、审核人、时间、版本", "一条任务一个当前决定"],
        ["Derived label", "review-apply", "原标签 + 已批准变更", "源标签不被覆盖"],
    ], [1900, 2000, 2900, 2560])

    chapter(doc, "38", "公开复现与故障注入实验", "用无 GPU fixture 和主动制造失败验证机制，而不是只看成功截图。")
    add_paragraph(doc, "公开演示分为数据审计、阈值/共识/近重复/主动审核、GT/AUTO 审核和安全写回四组。它们不声明模型精度，但能证明输入输出契约、决策规则、幂等和失败处理。")
    add_command(doc, r"""
    python examples\create_synthetic_dataset.py --output .demo-dataset
    yolo-label-recovery audit .demo-dataset --output-dir .demo-audit --hash-images --check-images
    """, "演示数据故意有问题。预期 FAIL 才表示审计工具正确捕获非法类别、孤立标签和跨划分重复。")
    add_command(doc, r"""
    python examples\create_review_fixture.py --output-dir .demo-review-fixture
    yolo-label-recovery review-build .demo-review-fixture\dataset .demo-review-fixture\candidates.csv `
      --output-dir .demo-review-result --render --redact-paths
    yolo-label-recovery review-ui .demo-review-result
    """, "审核 fixture 枚举 GT0_AUTO0、GT1_AUTO0、GT0_AUTO1、GT1_AUTO1；先观察允许动作，再做少量人工决定。")
    add_table(doc, ["故障实验", "操作", "预期系统行为", "证明的能力"], [
        ["重复导入", "同一 review_queue 执行两次", "第二次跳过已有 candidate_id", "幂等与唯一键"],
        ["旧版本提交", "保留旧页面后由另一端修改", "返回 409，不覆盖新决定", "乐观并发控制"],
        ["租约过期", "停止心跳并等待截止", "任务重新可领取", "故障回收"],
        ["路径穿越", "请求 review_root 外部路径", "403 拒绝", "资源级安全"],
        ["源 GT 漂移", "审核后修改待替换 GT", "review-apply 阻止替换", "审核依据一致性"],
        ["OOM", "提高 batch 直到显存不足", "未提交批次减半重试", "资源自适应与事务式提交"],
    ], [1700, 2700, 2700, 2260])
    add_callout(doc, "实验记录模板", "每次实验至少保存：初始状态、触发步骤、错误响应或日志、数据库/文件状态、恢复步骤和最终不变量。面试讲故障时，证据比‘我后来改好了’更有说服力。", "tip")

    chapter(doc, "39", "代码级面试：Spring Boot、MySQL 与 Vue 深挖", "把一次一键审核拆成 HTTP、事务、锁、租约、版本、审计和前端状态机。")
    add_paragraph(doc, "一次领取不是 Repository 查一行那么简单。Service 先检查项目访问权，再复用审核人仍有效的当前租约；没有活动任务时才在事务中用悲观锁查找第一条 PENDING 或过期 CLAIMED 任务，写入 claimed_by 和 lease_until，并追加审计事件。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/service/TaskService.java", 43, 84, "claimNext 的事务边界同时覆盖访问检查、活动任务复用、行锁领取、租约写入和审计。")
    add_code(doc, "platform/backend/src/main/java/com/jiapeng/labelreview/repository/TaskRepository.java", 18, 37, "PESSIMISTIC_WRITE 只保护短领取事务；查询条件同时允许 PENDING 和租约已过期的 CLAIMED。")
    add_paragraph(doc, "数据库约束提供最终防线：任务唯一键阻止重复候选，决定唯一键阻止一任务多决定，@Version 阻止旧页面覆盖，复合索引对应领取、本图候选和最近审核三条高频路径。")
    add_table(doc, ["机制", "时间尺度", "失败示例", "系统响应"], [
        ["悲观行锁", "一次领取事务", "两人同时点领取", "等待锁后拿到不同任务"],
        ["租约", "跨请求的分钟级占用", "浏览器断网或关闭", "到期后任务重新可领取"],
        ["乐观版本", "提交瞬间", "旧页面提交决定", "409 Conflict"],
        ["唯一键", "数据生命周期", "队列重复导入", "跳过或数据库拒绝重复"],
        ["审计事件", "事后追溯", "管理员修订决定", "保留操作者和前后动作"],
    ], [1700, 2000, 3000, 2660])
    add_paragraph(doc, "前端高频路径使用一键保存并自动前进：先提交当前决定，成功后优先找本图下一条 PENDING；没有则领取下一图。失败时保留当前任务和用户输入，不得假装已保存。")
    add_code(doc, "platform/frontend/src/App.vue", 119, 207, "claimNext、submitDecision 与 advanceAfterDecision 共同形成高吞吐但可恢复的前端状态机。")
    add_callout(doc, "高频追问", "为什么前端隐藏按钮不算授权？因为请求可以绕过界面。角色、项目范围、任务所有权、租约和版本必须由后端在每次请求中重新验证。", "warn")

    chapter(doc, "40", "用证据矩阵控制项目表述", "明确区分实现、真实验证、局部原型和架构设计，既不低估工作量，也不透支可信度。")
    add_table(doc, ["能力", "当前证据等级", "可以怎样说", "不能怎样说"], [
        ["Multi-Teacher 全量扫描", "实现 + 真实规模", "完成 174,426 次图片-模型推理", "Teacher 消除了全部漏标"],
        ["多人审核平台", "实现 + 双账号验证", "30,183 任务已导入并协作审核", "已经通过高并发生产压测"],
        ["历史决定迁移", "真实验证", "4,465 条决定幂等迁移", "CSV 本身就是数据库真值"],
        ["模型收益", "待固定集继续验证", "按六类和联合场景对比", "训练损失下降就说明更好"],
        ["20 路 RTSP", "系统设计", "已设计最新帧和动态 Batch 方案", "仓库已稳定承载 20 路"],
        ["RAG/Agent", "系统设计", "已定义证据与工具安全边界", "已自动执行真实高风险处置"],
    ], [2100, 1900, 3300, 2060])
    add_paragraph(doc, "真实数字必须带口径。29,071 张图片说明数据规模，174,426 次推理说明六个 Teacher 的计算工作量，30,183 条审核项说明人机协作任务规模；它们都不能直接推出 mAP、事件准确率或最大并发用户数。")
    add_formula(doc, "证据化表达", "结论强度 <= 最弱证据强度", "截图只能证明界面存在；测试证明特定行为；真实规模运行证明该范围可用；独立固定测试集才证明模型收益。")
    add_numbers(doc, [
        "先说业务问题与错误代价，不先报技术名词。",
        "再说关键取舍：单类 Teacher、源数据只读、人工门控和数据库并发。",
        "随后给真实数字和可复现 fixture，说明验证范围。",
        "最后主动说明 RTSP、RAG、Agent 和边缘部署的当前完成度与下一步。",
    ])
    add_callout(doc, "面试加分点", "能主动解释‘为什么这个数字不能证明另一个结论’，比堆更多漂亮指标更能体现工程判断。", "tip")

    chapter(doc, "41", "十四天把项目真正变成自己的能力", "用代码跟读、故障实验、二次开发和模拟面试完成从会用到会设计的跃迁。")
    add_table(doc, ["阶段", "天数", "必须完成", "可复核输出"], [
        ["问题与数据", "Day 1-2", "画出漏标错误监督和 GT/AUTO 四状态", "一页问题定义、三个手算框案例"],
        ["Python 链路", "Day 3-4", "跑公开 fixture，跟读 run/review/apply", "报告、候选 CSV、派生数据差异"],
        ["资源与恢复", "Day 5", "制造中断或 OOM，验证 resume", "前后 state、manifest 和去重证据"],
        ["后端与数据库", "Day 6-8", "双账号领取、过期、409、重复导入", "API 响应、数据库行和审计事件"],
        ["前端与部署", "Day 9-10", "修改一个审核交互并构建部署", "截图、构建产物和回归记录"],
        ["二次开发", "Day 11-12", "新增一个小能力与完整测试", "PR 风格说明和 CI 结果"],
        ["面试演练", "Day 13-14", "30 秒、3 分钟、30 分钟三版讲解", "录音复盘、追问清单和白板图"],
    ], [1700, 1100, 3600, 2960])
    add_paragraph(doc, "推荐的二次开发题目：增加项目级类别过滤；导出审核员效率报表；为历史修订增加原因枚举；为领取查询增加 Testcontainers MySQL 并发测试；为 review-apply 增加变更清单和回滚 manifest。每个题目都要同时修改契约、实现、测试、文档和演示。")
    add_table(doc, ["完成标准", "不合格表现", "合格表现"], [
        ["理解代码", "只能说用了哪些框架", "能追踪一次状态变化到 SQL 和失败响应"],
        ["理解算法", "背 IoU 公式", "能手算并解释 IoU/IoS/中心距离为何给出不同判断"],
        ["理解并发", "说用了锁", "能区分行锁、租约、版本和唯一键"],
        ["理解验证", "展示成功截图", "能设计失败实验并说明不变量"],
        ["理解边界", "把规划都说成已落地", "按证据等级准确表达"],
    ], [1900, 3300, 4160])
    add_callout(doc, "最终目标", "你应该能在没有这份手册时，从白板画出双闭环架构，打开仓库定位关键函数，构造一个失败场景，解释数据库如何保护状态，并明确哪些结论还需要下一轮评测。", "info")

    chapter(doc, "附录 A", "API 速查", "快速定位前后端契约。", new_page=False)
    add_table(doc, ["方法", "路径", "角色", "用途"], [
        ["POST", "/api/auth/login", "公开", "账号密码换 JWT"],
        ["GET", "/api/auth/me", "已登录", "当前用户"],
        ["POST", "/api/admin/users", "ADMIN", "创建协作账号"],
        ["GET/POST", "/api/projects", "成员/ADMIN", "列表或创建项目"],
        ["POST", "/api/projects/{id}/members", "ADMIN", "分配项目成员"],
        ["POST", "/api/projects/{id}/tasks:batch", "ADMIN", "幂等批量导入"],
        ["POST", "/api/projects/{id}/decisions:history", "ADMIN", "幂等迁移历史决定"],
        ["GET", "/api/projects/{id}/progress", "项目成员", "进度统计"],
        ["POST", "/api/tasks/claim-next", "REVIEWER/ADMIN", "原子领取"],
        ["POST", "/api/tasks/{id}/heartbeat", "任务所有者", "续租"],
        ["POST", "/api/tasks/{id}/release", "任务所有者", "释放"],
        ["POST", "/api/tasks/{id}/decision", "任务所有者", "提交决策"],
        ["PUT", "/api/tasks/{id}/decision", "原审核人/ADMIN", "修订已完成决定"],
        ["GET", "/api/tasks/recent?projectId={id}", "REVIEWER/ADMIN", "当前账号最近审核"],
        ["GET", "/api/tasks/{id}/image-candidates", "项目成员", "按原图聚合候选"],
        ["GET", "/api/tasks/{id}/visual", "项目成员", "读取受保护图片"],
    ], [1000, 3600, 1900, 2860])

    chapter(doc, "附录 B", "核心术语", "面试时用准确术语表达。")
    terms = [
        ["Teacher", "用于生成候选证据的专用检测模型，不等于最终真值。"],
        ["GT", "Ground Truth，已有人工标签。"],
        ["IoU", "交集除以并集，衡量两个框总体重叠。"],
        ["IoS", "交集除以较小框面积，识别包含关系。"],
        ["Lease", "带截止时间的临时任务所有权，需心跳续租。"],
        ["Pessimistic Lock", "读取时锁行，防止并发事务同时领取。"],
        ["Optimistic Version", "提交时比较版本，拒绝旧状态覆盖。"],
        ["Idempotency", "重复执行不会重复产生副作用。"],
        ["RBAC", "基于角色的权限控制，不能替代资源级授权。"],
        ["Flyway", "按版本管理数据库 schema 的迁移工具。"],
        ["Derived Dataset", "由源数据和审核决策生成的新数据集，源目录保持不变。"],
    ]
    add_table(doc, ["术语", "本项目中的含义"], terms, [2300, 7060])

    chapter(doc, "附录 C", "代码阅读清单与自测题", "用问题验证自己是否真正掌握。")
    add_bullets(doc, [
        "能否解释漏标为什么会在 BCE 分类损失中形成错误负监督，而不只是少一条训练样本？",
        "能否手算一个同目标大小框和小框包含大框案例的 IoU、IoS、归一化中心距离，并据此选择规则？",
        "能否解释 candidate_id 为什么必须跨 CSV、数据库、审核决定和写回日志保持稳定？",
        "能否指出 SecurityConfig 中哪些路径匿名、哪些路径按角色限制？",
        "能否解释 requireAccess 为什么必须存在于 Service，而不能只做前端隐藏？",
        "能否从 findClaimable 的查询条件解释租约过期任务如何回到队列？",
        "能否说明 @Version 与请求中的 expectedVersion 分别在哪一层生效？",
        "能否解释为什么 visual 路径要 normalize 后再 startsWith(root)？",
        "能否说明 CSV 导入为什么同时需要批内去重和数据库唯一键？",
        "能否解释历史决定迁移为何必须先建立完整任务，并复用 DecisionPolicy？",
        "能否设计一个双账号领取测试，区分行锁、租约、版本号与唯一键？",
        "能否在不看代码时写出 App.vue 的领取、心跳、提交、清理流程？",
        "能否说出当前未实现的至少五项生产化能力，并给出优先级？",
    ])
    add_callout(doc, "毕业标准", "你能独立完成 Lab 3 或 Lab 5，写出测试，并用 15 分钟讲清一次真实故障与修复，这个项目才真正属于你。", "tip")

    chapter(doc, "附录 D", "学习依赖与代码索引", "按‘先懂什么、再读哪里、最后怎么验证’建立可执行学习路线。")
    add_paragraph(doc, "下面的索引不是文件清单，而是知识依赖图。每一行都要求完成四步：先理解前置概念，再沿真实代码追踪正常路径，然后主动构造失败路径，最后用测试或数据库证据验证结论。")
    add_table(doc, ["主题", "前置知识", "代码入口", "验证任务"], [
        ["漏标与错误监督", "目标检测标签、BCE、正负样本分配", "单类数据准备脚本、Teacher 训练配置", "构造一张有 smoking 但无标签的图，解释其训练信号"],
        ["候选几何分类", "坐标系、IoU、IoS、中心距离、面积比", "geometry.py、case_engine.py", "手算三个框，再与单元测试输出对照"],
        ["流式多 Teacher 推理", "显存、内存、批处理、异常边界", "autolabel 流水线、state.py", "降低 batch 触发重试，验证断点和 run signature"],
        ["审核动作语义", "状态机、幂等、派生数据集", "review_apply.py、decision_policy.py", "分别执行 ADD、REPLACE、REJECT 并核对标签差异"],
        ["认证与授权", "HTTP、JWT、RBAC、资源所有权", "SecurityConfig、ProjectAccessService", "普通账号调用管理员接口和越权项目接口"],
        ["多人并发审核", "事务、行锁、租约、乐观锁、唯一键", "ReviewTaskService、Repository 查询", "双账号同时领取并用旧 version 提交"],
        ["前端审核状态", "异步请求、组件状态、Blob URL", "App.vue、API client", "快速切换任务，验证旧图片请求不会覆盖新任务"],
        ["交付与复现", "配置、迁移、日志、测试金字塔", "one-click 脚本、Flyway、CI", "空机器启动后完成登录、导入、审核、导出闭环"],
    ], [1800, 2200, 2450, 2910])
    add_callout(doc, "推荐学习顺序", "先完成算法链路：第 0、5、6、7、8、9、25、37 章；再完成系统链路：第 10 至 18、23、24、39 章；随后用第 38 章做故障实验，最后用第 36、40、41 章整理 GitHub 作品集和面试表达。每学完一段，都用第 25 章的一张图重新讲一遍。", "info")

    doc.add_heading("D.1 建议的七天学习节奏", level=2)
    add_table(doc, ["天", "学习目标", "必须动手", "当天输出"], [
        ["Day 1", "看懂问题、数据契约和漏标危害", "检查一组 YOLO 标签并画框", "一页问题定义和标签规范"],
        ["Day 2", "掌握 GT/AUTO 状态和多几何量", "手算并运行 geometry 测试", "四类状态和边界案例表"],
        ["Day 3", "掌握流式推理、阈值和恢复机制", "运行小规模 dry-run 与中断恢复", "资源曲线和恢复证据"],
        ["Day 4", "掌握审核动作和安全写回", "审核 10 个候选并生成派生集", "源/派生标签差异报告"],
        ["Day 5", "掌握 Spring Security、事务和并发", "双账号领取、过期、冲突实验", "接口时序和数据库状态截图"],
        ["Day 6", "掌握前端状态与局域网部署", "完成登录、图片加载和多人审核", "完整操作录屏或截图"],
        ["Day 7", "形成实验与面试表达", "固定测试集对比并模拟追问", "15 分钟项目讲稿和风险清单"],
    ], [900, 2700, 3200, 2560])
    add_callout(doc, "不要只读", "如果只阅读代码，你会产生‘好像懂了’的错觉。每个主题至少保留一种可复核证据：测试通过记录、数据库行、日志片段、可视化图片、接口响应或前后版本指标。", "warn")

    doc.add_heading("官方学习入口", level=2)
    sources = [
        ("Spring Boot", "https://spring.io/projects/spring-boot/"),
        ("Spring Security JWT Resource Server", "https://docs.spring.io/spring-security/reference/servlet/oauth2/resource-server/jwt.html"),
        ("Spring Security Password Storage", "https://docs.spring.io/spring-security/reference/features/authentication/password-storage.html"),
        ("Vue 3 Guide", "https://vuejs.org/guide/introduction.html"),
        ("Docker Compose", "https://docs.docker.com/compose/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        hyperlink(p, label, url)

    chapter(doc, "附录 E", "现场演示与面试速查", "在联调、答辩或面试前十分钟快速确认环境、入口、证据和边界。")
    add_table(doc, ["场景", "先打开/执行", "必须确认"], [
        ["GitHub 展示", "README.md -> README.zh-CN.md -> 文档总入口", "英文默认首页、中文深度手册、真实截图、CI 和相对链接正常"],
        ["无 GPU 演示", "create_synthetic_dataset + audit", "预期 FAIL 且问题数量符合 fixture"],
        ["审核演示", "review-build + review-ui", "同图候选、允许动作、自动保存和恢复"],
        ["多人平台", "登录 -> 项目 -> 领取 -> 决策 -> 最近审核", "两个账号不重复领取、纠错有审计"],
        ["代码讲解", "cli.py -> main pipeline -> review -> TaskService", "每段都能说输入、不变量和失败路径"],
        ["模型汇报", "固定 test 与联合场景专项集", "分类别指标、阈值、模型版本和硬件口径"],
    ], [1700, 3600, 4060])
    add_command(doc, r"""
    python scripts\validate_docs.py
    python tests\run_smoke_tests.py
    pytest
    cd platform\backend
    mvn test
    cd ..\frontend
    npm ci
    npm run build
    """, "提交或现场演示前的最小验证链。根据当前机器是否安装 Java、Node 和完整 Python 开发依赖选择执行。")
    add_bullets(doc, [
        "30 秒开场：业务问题、核心闭环、真实任务规模。",
        "3 分钟主线：错误监督、Teacher 证据、多几何审核、资源控制、多人一致性、固定评测。",
        "30 分钟深挖：代码入口、事务边界、失败实验、数据库约束、前端状态和演进边界。",
        "不要把截图当精度证据，不要把双账号试用称为高并发，不要把 RTSP/RAG/Agent 规划称为仓库已实现。",
    ])
    add_callout(doc, "最后自检", "如果面试官随机指向 README 的一句话，你都能在 30 秒内打开对应代码、测试或报告；如果不能，就先降低表述强度，再补证据。", "warn")

    # Metadata and save.
    props = doc.core_properties
    props.title = "YOLO Label Recovery 项目开发学习与大厂面试手册"
    props.subject = "多 Teacher YOLO 漏标恢复与多人审核协作平台"
    props.author = "YOLO Label Recovery Project Team"
    props.keywords = "YOLO, Multi-Teacher, Spring Boot, Vue, MySQL, JWT, RTSP, RAG, Agent, concurrency, interview"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the project development and interview handbook.")
    parser.add_argument("--output", type=Path, default=OUT, help="Destination .docx path")
    args = parser.parse_args()
    build_document(args.output)
