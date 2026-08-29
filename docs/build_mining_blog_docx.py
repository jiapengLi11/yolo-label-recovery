from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
MARKDOWN = ROOT / "MINING_SAFETY_AI_ENGINEERING_BLOG.zh-CN.md"
OUTPUT = ROOT / "矿区智能安全监控项目复盘博客_图文版.docx"
ASSET_DIR = ROOT / "assets"
GENERATED_DIR = ROOT / "_blog_docx_assets"


ACCENT = "0F766E"
INK = "132238"
MUTED = "64748B"
FILL = "ECFDF5"
CODE_FILL = "0F172A"
TABLE_FILL = "E2E8F0"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "CBD5E1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def style_run(run, size: float | None = None, bold: bool | None = None, color: str | None = None, font: str = "Microsoft YaHei") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_mixed_markdown(paragraph, text: str, size: float = 10.5, color: str = INK, bold_default: bool = False) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            style_run(run, size=size, bold=bold_default, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            style_run(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            style_run(run, size=size, bold=False, color="334155", font="Consolas")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        style_run(run, size=size, bold=bold_default, color=color)


def find_font() -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), 30)
    return ImageFont.load_default()


def diagram_flow(path: Path, title: str, nodes: Iterable[str], subtitle: str | None = None) -> None:
    font = find_font()
    small = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18) if Path(r"C:\Windows\Fonts\msyh.ttc").exists() else ImageFont.load_default()
    img = Image.new("RGB", (1500, 620), "white")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((30, 30, 1470, 590), radius=26, fill="#F8FAFC", outline="#CBD5E1", width=2)
    draw.text((70, 62), title, fill="#0F172A", font=font)
    if subtitle:
        draw.text((72, 108), subtitle, fill="#64748B", font=small)
    nodes = list(nodes)
    top_y = 210
    box_w = 235
    gap = (1320 - box_w * len(nodes)) // max(1, len(nodes) - 1)
    x = 90
    centers = []
    for i, label in enumerate(nodes):
        fill = "#ECFDF5" if i % 2 == 0 else "#EFF6FF"
        draw.rounded_rectangle((x, top_y, x + box_w, top_y + 150), radius=18, fill=fill, outline="#0F766E", width=3)
        wrapped = []
        line = ""
        for ch in label:
            if draw.textlength(line + ch, font=small) > box_w - 34:
                wrapped.append(line)
                line = ch
            else:
                line += ch
        if line:
            wrapped.append(line)
        yy = top_y + 42 - max(0, len(wrapped) - 2) * 12
        for w in wrapped:
            draw.text((x + 18, yy), w, fill="#134E4A", font=small)
            yy += 28
        centers.append((x + box_w, top_y + 75))
        if i < len(nodes) - 1:
            x2 = x + box_w + gap
            draw.line((x + box_w + 8, top_y + 75, x2 - 12, top_y + 75), fill="#0F766E", width=4)
            draw.polygon([(x2 - 12, top_y + 75), (x2 - 30, top_y + 65), (x2 - 30, top_y + 85)], fill="#0F766E")
        x += box_w + gap
    draw.text((72, 440), "关键原则：模型只提出证据，规则解释关系，人工授权变化，派生数据集负责可回滚。", fill="#334155", font=small)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def production_summary(path: Path) -> None:
    font = find_font()
    small = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 22) if Path(r"C:\Windows\Fonts\msyh.ttc").exists() else ImageFont.load_default()
    tiny = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18) if Path(r"C:\Windows\Fonts\msyh.ttc").exists() else ImageFont.load_default()
    img = Image.new("RGB", (1500, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((30, 30, 1470, 730), radius=26, fill="#F8FAFC", outline="#CBD5E1", width=2)
    draw.text((70, 70), "生产验证摘要", fill="#0F172A", font=font)
    cards = [
        ("30,183", "候选审核任务", "从六个 Teacher 全量扫描得到"),
        ("4,465", "历史决定迁移", "桌面审核结果幂等导入平台"),
        ("2", "账号并行验证", "校园网内多人同时领取审核"),
        ("6", "检测类别", "person / helmet / vest / tractor / slipper / smoking"),
    ]
    x, y = 80, 170
    for i, (num, label, desc) in enumerate(cards):
        cx = x + (i % 2) * 680
        cy = y + (i // 2) * 230
        draw.rounded_rectangle((cx, cy, cx + 620, cy + 170), radius=20, fill="#ECFDF5", outline="#0F766E", width=3)
        draw.text((cx + 34, cy + 30), num, fill="#0F766E", font=font)
        draw.text((cx + 34, cy + 82), label, fill="#0F172A", font=small)
        draw.text((cx + 34, cy + 122), desc, fill="#64748B", font=tiny)
    draw.text((82, 650), "注意：这些证据证明流程跑通，最终模型收益仍需固定测试集和联合场景专项集验证。", fill="#334155", font=tiny)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def resolve_image(src: str) -> Path:
    rel = Path(src)
    original = ROOT / rel
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    if original.suffix.lower() != ".svg":
        return original
    stem = original.stem
    out = GENERATED_DIR / f"{stem}.png"
    if stem == "mining-safety-ai-architecture":
        candidate = ROOT / "_handbook_assets" / "mining-system-architecture.png"
        if candidate.exists():
            return candidate
        diagram_flow(out, "矿区智能安全系统架构", ["设备与边缘层", "视频与推理层", "事件与业务层", "知识与智能层", "数据与模型闭环"])
    elif stem == "data-evidence-lifecycle":
        diagram_flow(out, "标签证据闭环", ["原始数据只读", "Teacher 扫描", "几何规则匹配", "人工审核授权", "派生数据集回写", "固定测试集评估"], "每一步都留下候选框、置信度、匹配关系和处理结果")
    elif stem == "production-validation-summary":
        production_summary(out)
    else:
        diagram_flow(out, stem, ["输入", "处理", "输出"])
    return out


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    style_run(run, size=9, color=MUTED)


def add_callout(doc: Document, text: str, title: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, FILL)
    set_cell_borders(cell, "99F6E4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    if title:
        r = p.add_run(title + "：")
        style_run(r, size=10.5, bold=True, color=ACCENT)
    add_mixed_markdown(p, text, size=10.5)
    doc.add_paragraph()


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size, color, before, after in [
        ("Heading 1", 17, ACCENT, 14, 8),
        ("Heading 2", 14, ACCENT, 12, 6),
        ("Heading 3", 12, "0F172A", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("YOLO Label Recovery | 技术复盘博客")
    style_run(r, size=8.5, color=MUTED)


def build_docx() -> Path:
    doc = Document()
    configure_doc(doc)
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    image_counter = 1
    pending_quote: list[str] = []

    def flush_code() -> None:
        if not code_lines:
            return
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_width(table)
        cell = table.cell(0, 0)
        set_cell_shading(cell, CODE_FILL)
        set_cell_borders(cell, "334155")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("\n".join(code_lines))
        style_run(run, size=9, color="E2E8F0", font="Consolas")
        doc.add_paragraph()
        code_lines.clear()

    def flush_quote() -> None:
        if pending_quote:
            add_callout(doc, " ".join(pending_quote), "核心观点")
            pending_quote.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_quote()
                in_code = True
                code_lines = []
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line:
            flush_quote()
            continue
        if line.startswith("> "):
            pending_quote.append(line[2:].strip())
            continue
        flush_quote()
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            caption = image_match.group(1).strip() or "项目配图"
            img_path = resolve_image(image_match.group(2).strip())
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                with Image.open(img_path) as im:
                    width_px, height_px = im.size
                width_in = 6.35
                if height_px > width_px * 0.8:
                    width_in = 4.8
                run.add_picture(str(img_path), width=Inches(width_in))
                add_caption(doc, f"图 {image_counter}  {caption}")
                image_counter += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[2:].strip())
            style_run(r, size=22, bold=True, color=INK)
            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = meta.add_run("图文版 | 适合复制到 CSDN、周报和面试复习")
            style_run(r, size=10, color=MUTED)
            doc.add_section(WD_SECTION_START.CONTINUOUS)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_mixed_markdown(p, re.sub(r"^\d+\.\s+", "", line), size=10.5)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_mixed_markdown(p, line[2:].strip(), size=10.5)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.25
        add_mixed_markdown(p, line, size=10.5)
    flush_quote()
    flush_code()
    doc.core_properties.title = "从漏标数据到矿区智能安全闭环"
    doc.core_properties.subject = "YOLO Label Recovery project case study"
    doc.core_properties.author = "李嘉鹏"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_docx())
